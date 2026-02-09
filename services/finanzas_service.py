from services.supabase_client import supabase
import os
from PySide6.QtWidgets import QMessageBox
import logging
from typing import Optional, Dict, Any
from decimal import Decimal, ROUND_HALF_UP
from datetime import date, datetime
import calendar
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo

# Configurar logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)



def obtener_ventas_ffmm_anual(anio: int):
    inicio = f"{anio}-01-01"
    fin = f"{anio}-12-31"

    response = (
        supabase
        .table("venta")
        .select("fecha_venta, codigo_interno, monto_venta, comprador_pago, vendedor_pago, tipo_documento")
        .eq("es_venta_proceso", False)
        .gte("fecha_venta", inicio)
        .lte("fecha_venta", fin)
        .execute()
    )

    return response.data or []


def obtener_arriendos_ffmm_anual(anio: int):
    inicio = f"{anio}-01-01"
    fin = f"{anio}-12-31"

    response = (
        supabase
        .table("arriendo")
        .select("id, codigo_interno, renta_mensual, fecha_inicio, fecha_termino, honorarios_porcentaje, cuenta_fm, tipo_documento, estado")
        .lte("fecha_inicio", fin)
        .or_(f"fecha_termino.is.null,fecha_termino.gte.{inicio}")
        .execute()
    )

    return response.data or []

def obtener_abonos_arriendo_anual(anio: int):
    inicio = f"{anio}-01-01"
    fin = f"{anio}-12-31"

    response = (
        supabase
        .table("abonos_arriendo")
        .select("arriendo_id, fecha, monto, descripcion")
        .gte("fecha", inicio)
        .lte("fecha", fin)
        .execute()
    )

    return response.data or []

def obtener_gastos_arriendo_anual(anio: int):
    inicio = f"{anio}-01-01"
    fin = f"{anio}-12-31"

    response = (
        supabase
        .table("gastos_arriendo")
        .select("arriendo_id, fecha, monto, descripcion")
        .gte("fecha", inicio)
        .lte("fecha", fin)
        .execute()
    )

    return response.data or []



def procesar_arriendos_anual(arriendos, anio):
    operaciones = []

    for a in arriendos:
        codigo = a["codigo_interno"]
        renta = Decimal(a["renta_mensual"])
        porcentaje_contrato = Decimal(a["honorarios_porcentaje"] or 0)
        cuenta_ffmm = a.get("cuenta_fm")
        tipo_documento = a.get("tipo_documento") or "Factura"
        estado = a.get("estado")

        fecha_inicio = date.fromisoformat(a["fecha_inicio"])

        if estado == "Vigente":
            fecha_termino = None
        else:
            fecha_termino = date.fromisoformat(a["fecha_termino"])

        inicio = max(fecha_inicio, date(anio, 1, 1))

        if fecha_termino:
            fin = min(fecha_termino, date(anio, 12, 31))
        else:
            fin = date(anio, 12, 31)

        primer_mes_contrato = date(fecha_inicio.year, fecha_inicio.month, 1)

        mes = inicio.month
        anio_iter = inicio.year

        while date(anio_iter, mes, 1) <= fin:

            fecha_mes_actual = date(anio_iter, mes, 1)

            if fecha_mes_actual == primer_mes_contrato:
                porcentaje = Decimal("50")
                tipo = "propietario"
            else:
                porcentaje = porcentaje_contrato
                tipo = "arrendatario"

            comision = renta * porcentaje / 100
            iva = comision * Decimal("19") / Decimal("119")
            honorarios_neto = comision - iva

            operaciones.append({
                "anio": anio_iter,
                "mes": mes,
                "codigo": codigo,
                "tipo": tipo,
                "base": renta,
                "porcentaje": porcentaje,
                "honorarios": honorarios_neto,
                "iva": iva,
                "total": comision,
                "cuenta_fm": cuenta_ffmm,
                "tipo_documento": tipo_documento
            })

            mes += 1
            if mes > 12:
                mes = 1
                anio_iter += 1

    return operaciones


def procesar_ventas_anual(ventas):
    operaciones = []

    for v in ventas:
        codigo = v["codigo_interno"]
        monto_venta = Decimal(v["monto_venta"])

        fecha = date.fromisoformat(v["fecha_venta"])
        anio = fecha.year
        mes = fecha.month
        comprador_pago = v["comprador_pago"]
        vendedor_pago = v["vendedor_pago"]
        tipo_documento = v["tipo_documento"]

        if comprador_pago == "Si" and vendedor_pago == "Si":
            # Comprador y vendedor
            for tipo in ("vendedor", "comprador"):
                porcentaje = Decimal("2")
                honorarios = monto_venta * porcentaje / 100
                iva = honorarios * Decimal("0.19")
                total = honorarios + iva

                operaciones.append({
                    "anio": anio,
                    "mes": mes,
                    "codigo": codigo,
                    "tipo": tipo,
                    "base": monto_venta,
                    "porcentaje": porcentaje,
                    "honorarios": honorarios,
                    "iva": iva,
                    "total": total,
                    "cuenta_fm":"",
                    "tipo_documento":tipo_documento
                })

    return operaciones


def procesar_abonos_arriendo(abonos, arriendos):
    operaciones = []

    # Mapeamos arriendo_id → arriendo completo
    arriendo_map = {a["id"]: a for a in arriendos}

    for ab in abonos:
        arriendo_id = ab["arriendo_id"]
        arriendo = arriendo_map.get(arriendo_id)
        if not arriendo:
            continue

        fecha = date.fromisoformat(ab["fecha"])
        monto = Decimal(ab["monto"])

        descripcion = ab["descripcion"]

        operaciones.append({
            "anio": fecha.year,
            "mes": fecha.month,
            "codigo": arriendo["codigo_interno"],
            "tipo": "arrendatario",
            "base": monto,
            "porcentaje": None,
            "honorarios": monto,
            "iva": 0,
            "total": monto,
            "cuenta_fm": arriendo.get("cuenta_fm", ""),
            "tipo_documento": "abono",
            "descripcion": descripcion
        })

    return operaciones

def procesar_gastos_arriendo(gastos, arriendos):
    operaciones = []

    # Mapeamos arriendo_id → arriendo completo
    arriendo_map = {a["id"]: a for a in arriendos}

    for ga in gastos:
        arriendo_id = ga["arriendo_id"]
        arriendo = arriendo_map.get(arriendo_id)
        if not arriendo:
            continue

        fecha = date.fromisoformat(ga["fecha"])
        monto = Decimal(ga["monto"])

        descripcion = ga["descripcion"]

        operaciones.append({
            "anio": fecha.year,
            "mes": fecha.month,
            "codigo": arriendo["codigo_interno"],
            "tipo": "propietario",
            "base": monto,
            "porcentaje": None,
            "honorarios": monto,
            "iva": 0,
            "total": monto,
            "cuenta_fm": arriendo.get("cuenta_fm", ""),
            "tipo_documento": "gasto",
            "descripcion": descripcion
        })

    return operaciones






def generar_excel_ffmm_anual(anio, mes_hasta, ruta_archivo):
    ventas = obtener_ventas_ffmm_anual(anio)
    arriendos = obtener_arriendos_ffmm_anual(anio)
    abonos = obtener_abonos_arriendo_anual(anio)
    gastos = obtener_gastos_arriendo_anual(anio)

    operaciones = []

    operaciones += procesar_arriendos_anual(arriendos, anio)
    operaciones += procesar_abonos_arriendo(abonos, arriendos)
    operaciones += procesar_gastos_arriendo(gastos, arriendos)
    operaciones += procesar_ventas_anual(ventas)

    # 🔹 FILTRO: desde enero hasta el mes seleccionado
    operaciones = [
        op for op in operaciones
        if op.get("anio") == anio and op.get("mes") is not None and op.get("mes") <= mes_hasta
    ]

    MESES_ES = {
        1: "Enero",
        2: "Febrero",
        3: "Marzo",
        4: "Abril",
        5: "Mayo",
        6: "Junio",
        7: "Julio",
        8: "Agosto",
        9: "Septiembre",
        10: "Octubre",
        11: "Noviembre",
        12: "Diciembre"
    }

    # Nombre del mes
    for op in operaciones:
        op["mes_nombre"] = MESES_ES.get(op["mes"], "")

    operaciones.sort(
        key=lambda x: (
            x["anio"],
            x["mes"],
            x["codigo"],
            x["tipo"]
        )
    )

    exportar_excel_ffmm(operaciones, ruta_archivo)


def exportar_excel_ffmm(operaciones, ruta_archivo):
    wb = Workbook()
    ws = wb.active
    ws.title = "FFMM"

    headers = [
        "Año\n",
        "Mes\n",
        "Mes\nNombre",
        "Código\nPropiedad",
        "Tipo",
        "Tipo\nDocumento",
        "Cuenta\nFFMM",
        "Monto",
        "Porcentaje",
        "Honorarios",
        "IVA",
        "Total",
        "Descripción"  # NUEVA COLUMNA
    ]

    ws.append(headers)

    # -------------------------------
    # Datos
    # -------------------------------
    for op in operaciones:
        ws.append([
            op.get("anio"),
            op.get("mes"),
            op.get("mes_nombre", ""),
            op.get("codigo"),
            op.get("tipo"),
            op.get("tipo_documento", ""),
            op.get("cuenta_fm", ""),
            float(op.get("base", 0)),
            float(op["porcentaje"]) / 100 if op.get("porcentaje") is not None else "",
            float(op.get("honorarios", 0)),
            float(op.get("iva", 0)),
            float(op.get("total", 0)),
            op.get("descripcion", "")  # NUEVA COLUMNA
        ])

    max_row = ws.max_row
    max_col = ws.max_column
    last_col_letter = ws.cell(row=1, column=max_col).column_letter

    # -------------------------------
    # Tabla Excel
    # -------------------------------
    tabla = Table(
        displayName="FFMM_Table",
        ref=f"A1:{last_col_letter}{max_row}"
    )

    tabla.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False
    )

    ws.add_table(tabla)

    # -------------------------------
    # Encabezados
    # -------------------------------
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

    ws.row_dimensions[1].height = 32

    # -------------------------------
    # Formatos
    # -------------------------------
    FORMATO_CLP = '"$"#,##0'
    FORMATO_PORCENTAJE = '0%'

    for row in ws.iter_rows(min_row=2, max_row=max_row):
        row[0].alignment = Alignment(horizontal="center")  # Año
        row[1].alignment = Alignment(horizontal="center")  # Mes
        row[4].alignment = Alignment(horizontal="center")  # Tipo
        row[5].alignment = Alignment(horizontal="center")  # Tipo Documento
        row[12].alignment = Alignment(horizontal="left", wrap_text=True)  # Descripción

        row[7].number_format = FORMATO_CLP        # Base
        row[8].number_format = FORMATO_PORCENTAJE # %
        row[9].number_format = FORMATO_CLP        # Honorarios
        row[10].number_format = FORMATO_CLP       # IVA
        row[11].number_format = FORMATO_CLP       # Total

    # -------------------------------
    # Totales acumulados (FUERA DE LA TABLA)
    # -------------------------------
    fila_totales = max_row + 2  # deja una fila en blanco

    total_honorarios = Decimal("0")
    total_iva = Decimal("0")
    total_general = Decimal("0")

    for op in operaciones:
        # SOLO operaciones contables (sin abonos ni gastos)
        if op.get("porcentaje") not in (None, "", 0):
            total_honorarios += Decimal(str(op["honorarios"]))
            total_iva += Decimal(str(op["iva"]))
            total_general += Decimal(str(op["total"]))

    # 🔴 redondeo FINAL (no antes)
    total_honorarios = total_honorarios.quantize(
        Decimal("1"), rounding=ROUND_HALF_UP
    )
    total_iva = total_iva.quantize(
        Decimal("1"), rounding=ROUND_HALF_UP
    )
    total_general = total_general.quantize(
        Decimal("1"), rounding=ROUND_HALF_UP
    )


    ws[f"I{fila_totales}"] = "TOTAL HONORARIOS"
    ws[f"J{fila_totales}"] = float(total_honorarios)

    ws[f"I{fila_totales + 1}"] = "TOTAL IVA"
    ws[f"J{fila_totales + 1}"] = float(total_iva)

    ws[f"I{fila_totales + 2}"] = "TOTAL GENERAL"
    ws[f"J{fila_totales + 2}"] = float(total_general)

    for f in range(fila_totales, fila_totales + 3):
        ws[f"I{f}"].font = Font(bold=True)
        ws[f"I{f}"].alignment = Alignment(
            horizontal="right",
            vertical="center",
            wrap_text=True
        )

        ws[f"J{f}"].font = Font(bold=True)
        ws[f"J{f}"].number_format = FORMATO_CLP
        ws[f"J{f}"].alignment = Alignment(
            horizontal="right",
            vertical="center"
        )

    # -------------------------------
    # Anchos de columnas
    # -------------------------------
    ANCHOS = {
        "A": 8,    # Año
        "B": 6,    # Mes
        "C": 14,   # Mes Nombre
        "D": 20,   # Código Propiedad
        "E": 14,   # Tipo
        "F": 18,   # Tipo Documento
        "G": 16,   # Cuenta FFMM
        "H": 16,   # Monto
        "I": 12,   # Porcentaje
        "J": 16,   # Honorarios
        "K": 14,   # IVA
        "L": 16,   # Total
        "M": 30    # Descripción (NUEVA)
    }

    for col, width in ANCHOS.items():
        ws.column_dimensions[col].width = width

    # -------------------------------
    # Encabezado fijo
    # -------------------------------
    ws.freeze_panes = "A2"

    wb.save(ruta_archivo)