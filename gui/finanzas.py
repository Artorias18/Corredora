from PySide6.QtWidgets import QHeaderView, QFileDialog, QMainWindow,QHBoxLayout,QTextEdit,QTabWidget,QDoubleSpinBox,QHeaderView,QFrame,QWidget,QLabel, QVBoxLayout,QAbstractItemView,QTableWidget,QTableWidgetItem,QScrollArea, QFormLayout, QLineEdit, QComboBox, QPushButton,QLabel, QDateEdit, QCheckBox, QMessageBox, QTableWidget, QTableWidgetItem,QSpinBox, QDialog
from PySide6.QtCore import Qt, QDate, Signal, QTimer
# from services.finanzas_service import obtener_arriendos_finanzas
from PySide6.QtGui import QIntValidator, QColor, QDoubleValidator
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows
from gui.usuario_actual import UsuarioActual
from services.supabase_client import supabase
import pandas as pd
import os

class DashboardFinanzas(QWidget):
    def __init__(self):
        super().__init__()
        layout = QVBoxLayout()
        self.setLayout(layout)

        # Tabla de arriendos
        self.tabla_arriendos = QTableWidget()
        self.tabla_arriendos.setColumnCount(7)
        self.tabla_arriendos.setHorizontalHeaderLabels([
            "ROL",
            "Destino",
            "DFL2",
            "Naturaleza bien raiz",
            "Comuna",
            "Rut dueño",
            "Rut arrendatario"
        ])
        self.tabla_arriendos.setSortingEnabled(True)
        self.tabla_arriendos.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.tabla_arriendos.setEditTriggers(QAbstractItemView.DoubleClicked)
        self.tabla_arriendos.horizontalHeader().setSectionResizeMode(QHeaderView.Interactive)
        self.tabla_arriendos.verticalHeader().setVisible(False)

        layout.addWidget(self.tabla_arriendos)

        # Botón exportar
        button_layout = QHBoxLayout()
        self.btn_exportar = QPushButton("Exportar a Excel")
        self.btn_exportar.clicked.connect(self.exportar_excel)
        button_layout.addWidget(self.btn_exportar)
        layout.addLayout(button_layout)


    #### ---------------- FUNCIONES DE EXCEL ---------------- ####

    def exportar_excel(operaciones, ruta):
        wb = Workbook()
        ws = wb.active
        ws.title = "Honorarios FFMM"

        ws.append([
            "Año",
            "Mes",
            "Propiedad",
            "Tipo Honorario",
            "Base",
            "%",
            "Honorarios",
            "IVA",
            "Monto FFMM"
        ])

        for op in operaciones:
            ws.append([
                op["anio"],
                op["mes"],
                op["codigo"],
                op["tipo"],
                float(op["base"]),
                float(op["porcentaje"]) if op["porcentaje"] else "",
                float(op["honorarios"]),
                float(op["iva"]),
                float(op["monto_ffmm"])
            ])

        wb.save(ruta)

