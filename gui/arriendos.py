from PySide6.QtWidgets import QHeaderView,QGroupBox ,QDialogButtonBox,QSizePolicy, QFileDialog, QMainWindow,QHBoxLayout,QTextEdit,QTabWidget,QDoubleSpinBox,QHeaderView,QFrame,QWidget,QLabel, QVBoxLayout,QAbstractItemView,QTableWidget,QTableWidgetItem,QScrollArea, QFormLayout, QLineEdit, QComboBox, QPushButton,QLabel, QDateEdit, QCheckBox, QMessageBox, QTableWidget, QTableWidgetItem,QSpinBox, QDialog
from services.arriendos_service import obtener_arriendos_resumen,CASTIGO_DEFAULT,IVA_FACTOR,calcular_evaluacion_independiente,safe_numeric, eliminar_arriendo, guardar_gastos_arriendo,obtener_ultima_evaluacion_arrendatario,guardar_abonos_arriendo, obtener_detalle_arriendo,obtener_arriendos_finanzas,obtener_arriendos_por_estado, guardar_arriendo
from PySide6.QtCore import Qt, QDate, Signal, QTimer
from PySide6.QtGui import QIntValidator, QColor, QDoubleValidator, QFont
from gui.usuario_actual import UsuarioActual
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows
from services.supabase_client import supabase
import json
import calendar
from datetime import date
from calendar import monthrange
import pandas as pd


class DetalleArriendoWindow(QWidget):
    def __init__(self, arriendo_id, dashboard = None):
        super().__init__()

        self.arriendo_id = arriendo_id

        self.dashboard = dashboard

        self.setWindowTitle(f"Detalle de Arriendo #{arriendo_id}")
        self.resize(900, 700)
        
        # Layout principal con scroll
        layout_principal = QVBoxLayout(self)
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        contenido = QWidget()
        scroll.setWidget(contenido)
        layout_principal.addWidget(scroll)

        # Layout del contenido
        layout_contenido = QVBoxLayout(contenido)


        self.detalle_arriendo = obtener_detalle_arriendo(arriendo_id)

        if not self.detalle_arriendo:
            layout_contenido.addWidget(QLabel("No se encontraron detalles para esta venta"))
            return
        
        self.tabs = QTabWidget()
        layout_contenido.addWidget(self.tabs)
        

        self.setup_tab_arrendador()
       
        self.setup_tab_arrendatario()

        self.setup_tab_propiedad()
       
        
    
        self.setup_tab_arriendo()
      








         # Botón para cerrar
        btn_cerrar = DoubleClickButton("Cerrar")
        btn_cerrar.doubleClicked.connect(self.close)
        layout_contenido.addWidget(btn_cerrar)
        
        self.btn_editar = DoubleClickButton("Editar Arriendo")
        self.btn_editar.doubleClicked.connect(self.abrir_formulario_edicion)
        layout_contenido.addWidget(self.btn_editar)

        self.btnEliminar = DoubleClickButton("Eliminar Arriendo")
        self.btnEliminar.doubleClicked.connect(self.eliminar_arriendo)
        layout_contenido.addWidget(self.btnEliminar)


    def setup_tab_arrendador(self):
        tab = QWidget()
        self.tabs.addTab(tab,"Arrendador")
        layout = QFormLayout(tab)

        arrendador = self.detalle_arriendo.get('arrendador', {})

        campos = [
            ("Nombre", arrendador.get('nombre')),
            ("Rut", arrendador.get('rut')),
            ("Teléfono", arrendador.get('telefono')),
            ("Dirección", arrendador.get('direccion')),
            ('Correo', arrendador.get('correo_electronico'))
        ]

        for label, value in campos:
            if value:
                layout.addRow(QLabel(f"<b>{label}:</b>"), QLabel(str(value)))
        


    
    def setup_tab_arrendatario(self):
        tab = QWidget()
        self.tabs.addTab(tab,"Arrendatario")
        layout = QFormLayout(tab)

        arrendatario = self.detalle_arriendo.get('arrendatario', {})

        campos = [
            ("Nombre:", arrendatario.get('nombre')),
            ("Rut:", arrendatario.get('rut')),
            ("Teléfono:", arrendatario.get('telefono')),
            ("Dirección:", arrendatario.get('direccion')),
            ('Correo:', arrendatario.get('email')),
            ("Tipo de Trabajador:", arrendatario.get('tipo_trabajador')),
            ("Antigüedad Laboral:", arrendatario.get('antiguedad_laboral')),
            ("¿Tiene Dicom?", arrendatario.get('dicom')),
            ("Estado de Evaluación", arrendatario.get('evaluacion_estado')),
            ("Renta:", f"${int(arrendatario.get('renta')):,}".replace(",", ".")),
            ("Comentarios:", arrendatario.get('comentarios'))
            ]
        
        for label, value in campos:
            if value:
                layout.addRow(QLabel(f"<b>{label}:</b>"), QLabel(str(value)))

    def setup_tab_propiedad(self):
        tab = QWidget()
        self.tabs.addTab(tab, "Propiedad")
        layout = QFormLayout(tab)
        
        propiedad = self.detalle_arriendo.get('propiedad', {})

        campos= [
            ("Código interno:", propiedad.get('codigo_interno')),
            ("Dirección:", propiedad.get('direccion')),
            ("ROL:", propiedad.get('rol')),
            ("Comuna:", propiedad.get('comuna')),
            ("Dominio vigente:", propiedad.get('dominio_vigente'))
        ]

        for label, value in campos:
            if value:
                layout.addRow(QLabel(f"<b>{label}:</b>"), QLabel(str(value)))

    def setup_tab_arriendo(self):
        tab = QWidget()
        self.tabs.addTab(tab, "Arriendo")

        # =========================
        # Layout principal del tab
        # =========================
        main_layout = QVBoxLayout(tab)
        main_layout.setContentsMargins(0, 0, 0, 0)
        main_layout.setSpacing(0)

        arriendo = self.detalle_arriendo.get('arriendo', {})

        # =========================
        # Content widget (dentro del scroll)
        # =========================
        content_widget = QWidget()
        content_vbox = QVBoxLayout(content_widget)
        content_vbox.setContentsMargins(8, 8, 8, 8)
        content_vbox.setSpacing(10)
        content_vbox.setAlignment(Qt.AlignTop)

        # ======================================
        # SECCIÓN 1: DATOS BÁSICOS
        # ======================================
        box_basicos = QGroupBox("Datos Básicos del Arriendo")
        layout_basicos = QFormLayout()

        campos_basicos = [
            ("Fecha Inicio", arriendo.get("fecha_inicio")),
            ("Fecha de Término", arriendo.get("fecha_termino")),
            ("Renta Mensual", self.format_money(arriendo.get('renta_mensual'))),
            ("Garantía", self.format_money(arriendo.get('garantia'))),
            ("Gastos comunes incluidos", "Sí" if arriendo.get("gastos_comunes_incluidos") else "No"),
            ("Estado del arriendo", arriendo.get("estado")),
            ("Tipo de contrato", arriendo.get("tipo_contrato")),
        ]

        for label, value in campos_basicos:
            if value not in (None, "", 0):
                layout_basicos.addRow(QLabel(f"<b>{label}:</b>"), QLabel(str(value)))

        box_basicos.setLayout(layout_basicos)
        box_basicos.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Maximum)
        content_vbox.addWidget(box_basicos)

        # ======================================
        # SECCIÓN 2: PAGOS
        # ======================================
        box_pagos = QGroupBox("Datos de Pago")
        layout_pagos = QFormLayout()

        campos_pagos = [
            ("Forma de pago", arriendo.get("forma_pago")),
            ("Periodo de pago", arriendo.get("periodo_pago")),
            ("Día de pago", arriendo.get("dia_pago")),
            ("Cuenta FM", arriendo.get("cuenta_fm")),
        ]

        for label, value in campos_pagos:
            if value not in (None, "", 0):
                layout_pagos.addRow(QLabel(f"<b>{label}:</b>"), QLabel(str(value)))

        box_pagos.setLayout(layout_pagos)
        box_pagos.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Maximum)
        content_vbox.addWidget(box_pagos)

        # ======================================
        # SECCIÓN 3: DEPÓSITOS
        # ======================================
        box_depositos = QGroupBox("Depósitos / Cuenta destino")
        layout_depositos = QFormLayout()

        campos_depositos = [
            ("Número de cuenta", arriendo.get("nro_cuenta")),
            ("Banco destino", arriendo.get("banco_destino")),
            ("Tipo de cuenta", arriendo.get("tipo_cuenta")),
            ("RUT para depósito", arriendo.get("rut_para_deposito")),
            ("Titular depósito", arriendo.get("titular_deposito")),
            ("Quién deposita", arriendo.get("quien_deposita")),
            ("Correo depósito", arriendo.get("correo_deposito")),
        ]

        for label, value in campos_depositos:
            if value not in (None, "", 0):
                layout_depositos.addRow(QLabel(f"<b>{label}:</b>"), QLabel(str(value)))

        box_depositos.setLayout(layout_depositos)
        box_depositos.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Maximum)
        content_vbox.addWidget(box_depositos)

        # ======================================
        # SECCIÓN 4: HONORARIOS
        # ======================================
        box_honorarios = QGroupBox("Honorarios")
        layout_honorarios = QFormLayout()

        campos_honorarios = [
            ("Honorarios (%):", f"{float((arriendo.get('honorarios_porcentaje') or 0) * 100):.1f}%"),
            ("Honorarios (monto):", self.format_money(arriendo.get('honorarios_monto')))
        ]

        for label, value in campos_honorarios:
            if value not in (None, "", 0):
                layout_honorarios.addRow(QLabel(f"<b>{label}:</b>"), QLabel(str(value)))

        box_honorarios.setLayout(layout_honorarios)
        box_honorarios.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Maximum)
        content_vbox.addWidget(box_honorarios)

        # ======================================
        # SECCIÓN 5: INFORMACIÓN ADICIONAL
        # ======================================
        box_extra = QGroupBox("Información adicional")
        layout_extra = QFormLayout()

        campos_extra = [
            ("Último mes pago:", arriendo.get("ultimo_mes_pago")),
            ("Aseo municipal:", arriendo.get("aseo_municipal")),
            ("Reajuste:", self.format_money(arriendo.get('reajuste'))),
            ("GGCC:", arriendo.get("ggcc")),
            ("Cuenta GGCC:", arriendo.get("cuenta_ggcc")),
            ("Dirección consulta:", arriendo.get("direccion_consulta")),
            ("Periodo anterior:", arriendo.get("periodo_anterior")),
            ("Monto anterior:", self.format_money(arriendo.get('monto_anterior'))),
            ("Naturaleza bien raíz:", arriendo.get("naturaleza_bien_raiz")),
            ("DFL12:", arriendo.get("dfl12")),
            ("Destino:", arriendo.get("destino")),
            ("Amoblado:", arriendo.get("amoblado")),
            ("Observaciones:", arriendo.get("observaciones")),
        ]

        for label, value in campos_extra:
            if value not in (None, "", 0):
                layout_extra.addRow(QLabel(f"<b>{label}:</b>"), QLabel(str(value)))

        box_extra.setLayout(layout_extra)
        box_extra.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Maximum)
        content_vbox.addWidget(box_extra)

        # Empuja todo hacia arriba
        content_vbox.addStretch()

        # =========================
        # Scroll Area
        # =========================
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setFrameStyle(0)
        scroll.setWidget(content_widget)
        scroll.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)

        # =========================
        # Añadir al tab
        # =========================
        main_layout.addWidget(scroll, 1)


    def abrir_formulario_edicion(self):
        
        try:
            arriendo = self.detalle_arriendo.get('arriendo', {})

            self.formulario = FormularioArriendo(arriendo_id= self.arriendo_id)

            self.formulario.cargar_datos_arriendo(arriendo_id = self.arriendo_id)

            self.formulario.arriendo_guardado.connect(self.dashboard.cargar_arriendos)

            self.formulario.show()

            self.close()

        except Exception as e:
            QMessageBox.warning(self, "Error", f"No se pudo abrir el formulario: {e}")

    def format_money(self, value):
        num = safe_numeric(value)
        return f"${int(num):,}".replace(",", ".")

    
    def eliminar_arriendo(self):
        try:
            # 1️⃣ Obtener el código interno de la venta
            codigo_interno = self.detalle_arriendo.get("propiedad", {}).get("codigo_interno")
            
            if not codigo_interno:
                QMessageBox.warning(self, "Eliminar arriendo", "No se encontró el código interno de este arriendo.")
                return

            # 2️⃣ Confirmar con el usuario
            confirm = QMessageBox.question(
                self,
                "Confirmar eliminación",
                f"¿Seguro que deseas eliminar el arriendo con código '{codigo_interno}'?\n"
                "Esto eliminará todos los registros.",
                QMessageBox.Yes | QMessageBox.No
            )

            if confirm != QMessageBox.Yes:
                return

            # 3️⃣ Llamar al backend
            resultado = eliminar_arriendo(codigo_interno)

            QMessageBox.information(self, "Eliminar arriendo", str(resultado))

            if self.dashboard:
                try:
                    self.dashboard.cargar_arriendos()
                except:
                    pass
            self.close()
        except Exception as e:
            QMessageBox.critical(self, "Error", f"Ocurrió un error al eliminar el arriendo:\n{str(e)}")
    

class DashboardArriendos(QWidget):
    def __init__(self,arriendo_id=None, parent = None):
        super().__init__(parent)
        layout = QVBoxLayout(self)
        self.cargando_tabla = False
        self.arriendo_id = arriendo_id
        self.detalle_arriendo = {}
        filter_layout = QHBoxLayout()
        
        # Filtro de búsqueda
        self.search_input = QLineEdit()
        self.search_input.setPlaceholderText("Buscar por arrendador, arrendatario o propiedad...")
        self.search_input.textChanged.connect(self.filtrar_arriendos)


         # Filtro de estado
        self.filtro_estado = QComboBox()
        self.filtro_estado.addItems([
            "Todos los arriendos",
            "Vigentes",
            "Finalizados",
            "Rescindidos"
            
        ])
        self.filtro_estado.currentIndexChanged.connect(self.filtrar_arriendos)
        
        # Filtro de fecha
        self.filtro_fecha = QComboBox()
        self.filtro_fecha.addItems([
            "Todas las fechas",
            "Últimos 7 días",
            "Últimos 30 días",
            "Este mes",
            "Este año"
        ])
        self.filtro_fecha.currentIndexChanged.connect(self.filtrar_arriendos)

        self.cmb_mes_desde = QComboBox()
        self.cmb_mes_hasta = QComboBox()

        MESES = [
            (1, "Enero"), (2, "Febrero"), (3, "Marzo"), (4, "Abril"),
            (5, "Mayo"), (6, "Junio"), (7, "Julio"), (8, "Agosto"),
            (9, "Septiembre"), (10, "Octubre"), (11, "Noviembre"), (12, "Diciembre")
        ]

        for numero, nombre in MESES:
            self.cmb_mes_desde.addItem(nombre, numero)
            self.cmb_mes_hasta.addItem(nombre, numero)

        # Valores por defecto (todo el año)
        self.cmb_mes_desde.setCurrentIndex(0)
        self.cmb_mes_hasta.setCurrentIndex(11)


        # Agregar filtros al layout
        filter_layout.addWidget(QLabel("Filtros:"))
        filter_layout.addWidget(self.search_input)
        filter_layout.addWidget(QLabel("Estado:"))
        filter_layout.addWidget(self.filtro_estado)
        filter_layout.addWidget(QLabel("Fecha:"))
        filter_layout.addWidget(self.filtro_fecha)

        layout.addLayout(filter_layout)
        filter_layout.addWidget(QLabel("Mes desde:"))
        filter_layout.addWidget(self.cmb_mes_desde)

        filter_layout.addWidget(QLabel("Mes hasta:"))
        filter_layout.addWidget(self.cmb_mes_hasta)

        btn_abonos = DoubleClickButton("Agregar abonos")
        btn_gastos = DoubleClickButton("Agregar gastos")

        btn_abonos.setEnabled(False)
        btn_gastos.setEnabled(False)

        btn_abonos.doubleClicked.connect(self.abrir_abonos_arriendo)
        btn_gastos.doubleClicked.connect(self.abrir_gastos_arriendo)

        botones_layout = QHBoxLayout()
        botones_layout.addWidget(btn_abonos)
        botones_layout.addWidget(btn_gastos)

        layout.addLayout(botones_layout)

        self.btn_abonos = btn_abonos
        self.btn_gastos = btn_gastos


         # Tabla de arriendos
        self.tabla_arriendos = QTableWidget()
        self.tabla_arriendos.setColumnCount(7)
        self.tabla_arriendos.setHorizontalHeaderLabels([
            "ID", 
            "Arrendador", 
            "Arrendatario", 
            "Propiedad", 
            "Fecha inicio",
            "Estado",
            "Tipo de contrato"
        ])
        self.tabla_arriendos.setSortingEnabled(True)
        self.tabla_arriendos.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.tabla_arriendos.setEditTriggers(QAbstractItemView.DoubleClicked)
        self.tabla_arriendos.horizontalHeader().setSectionResizeMode(QHeaderView.Interactive)
        self.tabla_arriendos.verticalHeader().setVisible(False)
        self.tabla_arriendos.itemDoubleClicked.connect(self.mostrar_detalle_arriendo)
        self.tabla_arriendos.itemSelectionChanged.connect(self.habilitar_boton_abonos)
        self.tabla_arriendos.itemSelectionChanged.connect(self.habilitar_boton_gastos)
        
        layout.addWidget(self.tabla_arriendos)

        self.tabla_arriendos.setStyleSheet("""
            QTableWidget::item:selected {
                background-color: #D5D8DC;
                color: black;
            }
        """)


        self.tabla_arriendos.resizeColumnsToContents()

        button_layout = QHBoxLayout()

        
        self.btn_agregar_arriendo = DoubleClickButton("Agregar Arriendo")
        self.btn_agregar_arriendo.doubleClicked.connect(self.abrir_formulario_arriendo)
        
        self.btn_actualizar = QPushButton("Actualizar Lista")
        self.btn_actualizar.clicked.connect(self.cargar_arriendos)
        
        self.btn_exportar = DoubleClickButton("Exportar a Excel")
        self.btn_exportar.doubleClicked.connect(self.exportar_excel)
        
        self.btn_agregar_factor_actualizacion = DoubleClickButton("Agregar Factor Actualización")
        self.btn_agregar_factor_actualizacion.doubleClicked.connect(self.abrir_factor_actualizacion)
       

 
        
        button_layout.addWidget(self.btn_agregar_arriendo)
        button_layout.addWidget(self.btn_actualizar)
        button_layout.addWidget(self.btn_agregar_factor_actualizacion)
        button_layout.addWidget(self.btn_exportar)
        
        
        layout.addLayout(button_layout)
        
        if self.arriendo_id:
            self.detalle_arriendo = obtener_detalle_arriendo(self.arriendo_id)
            if not self.detalle_arriendo:
                layout.addWidget(QLabel("No se encontraron detalles para este arriendo"))
            else:
                self.cargar_datos_arriendo(self.arriendo_id)



        # Cargar datos iniciales
        self.cargar_arriendos()

    def cargar_arriendos(self):
        try:
            self.cargando_tabla = True

            # 🔒 BLOQUEO TOTAL (CLAVE)
            self.tabla_arriendos.setSortingEnabled(False)
            self.tabla_arriendos.blockSignals(True)
            self.tabla_arriendos.setUpdatesEnabled(False)

            self.tabla_arriendos.clearContents()
            self.tabla_arriendos.setRowCount(0)

            estado_filtro = self.filtro_estado.currentText()

            if estado_filtro == "Todos los arriendos":
                arriendos = obtener_arriendos_resumen()
            else:
                arriendos = obtener_arriendos_por_estado(estado_filtro)

            self.arriendos = arriendos or []

            if not arriendos:
                return

            self.tabla_arriendos.setRowCount(len(arriendos))

            for row, arriendo in enumerate(arriendos):
                self.tabla_arriendos.setItem(row, 0, QTableWidgetItem(str(arriendo.get("id", ""))))
                self.tabla_arriendos.setItem(row, 1, QTableWidgetItem(arriendo.get("arrendador", "")))
                self.tabla_arriendos.setItem(row, 2, QTableWidgetItem(arriendo.get("arrendatario", "")))
                self.tabla_arriendos.setItem(row, 3, QTableWidgetItem(arriendo.get("propiedad", "")))
                self.tabla_arriendos.setItem(
                    row, 4, QTableWidgetItem(str(arriendo.get("fecha_inicio", "")))
                )
                self.tabla_arriendos.setItem(
                    row, 5, QTableWidgetItem(arriendo.get("estado", ""))
                )
                self.tabla_arriendos.setItem(
                    row, 6, QTableWidgetItem(arriendo.get("tipo_contrato", ""))
                )

            # 🎯 AJUSTES VISUALES (AÚN SIN SORTING)
            self.tabla_arriendos.resizeColumnsToContents()
            self.tabla_arriendos.setColumnHidden(0, True)

        except Exception as e:
            QMessageBox.critical(
                self, "Error", f"No se pudieron cargar los arriendos: {str(e)}"
            )

        finally:
            # 🔓 REACTIVAR TODO
            self.tabla_arriendos.setUpdatesEnabled(True)
            self.tabla_arriendos.blockSignals(False)
            self.tabla_arriendos.setSortingEnabled(True)

            self.cargando_tabla = False



    def filtrar_arriendos(self):
        texto = self.search_input.text().lower()
        estado = self.filtro_estado.currentText()
        fecha = self.filtro_fecha.currentText()

        for row in range(self.tabla_arriendos.rowCount()):
            mostrar_fila = True
            
            # Filtrar por texto
            if texto:
                mostrar_fila = any(
                    texto in self.tabla_arriendos.item(row, col).text().lower()
                    for col in range(self.tabla_arriendos.columnCount())
                )
            
            # Filtrar por estado
            if mostrar_fila and estado != "Todos los arriendos":
                estado_item = self.tabla_arriendos.item(row, 5).text()
                
                if estado == "Vigentes":
                    mostrar_fila = estado_item == "Vigente"
                elif estado == "Finalizados":
                    mostrar_fila = estado_item == "Finalizado"

                elif estado == "Rescindidos":
                    mostrar_fila = estado_item == "Rescindido"

            
            # Filtrar por fecha
            if mostrar_fila and fecha != "Todas las fechas":
                fecha_item = self.tabla_arriendos.item(row, 4).text()
                fecha_arriendo = QDate.fromString(fecha_item, "yyyy-MM-dd")
                hoy = QDate.currentDate()
                
                if fecha == "Últimos 7 días":
                    mostrar_fila = fecha_arriendo.daysTo(hoy) <= 7
                elif fecha == "Últimos 30 días":
                    mostrar_fila = fecha_arriendo.daysTo(hoy) <= 30
                elif fecha == "Este mes":
                    mostrar_fila = fecha_arriendo.month() == hoy.month() and fecha_arriendo.year() == hoy.year()
                elif fecha == "Este año":
                    mostrar_fila = fecha_arriendo.year() == hoy.year()
            
            self.tabla_arriendos.setRowHidden(row, not mostrar_fila)

    
        
    def mostrar_detalle_arriendo(self, item):
        arriendo_id = int(self.tabla_arriendos.item(item.row(), 0).text())
        self.ventana_detalle = DetalleArriendoWindow(arriendo_id, dashboard=self)
        self.ventana_detalle.show()

    def habilitar_boton_abonos(self):
        self.btn_abonos.setEnabled(
            len(self.tabla_arriendos.selectedItems()) > 0
        )

    def habilitar_boton_gastos(self):
        self.btn_gastos.setEnabled(
            len(self.tabla_arriendos.selectedItems()) > 0
        )
    
    def obtener_arriendo_seleccionado(self):
        fila = self.tabla_arriendos.currentRow()
        if fila < 0:
            return None

        arriendo_id = int(self.tabla_arriendos.item(fila, 0).text())
        return arriendo_id

    def abrir_abonos_arriendo(self):
        arriendo_id = self.obtener_arriendo_seleccionado()
        if not arriendo_id:
            QMessageBox.warning(self, "Error", "Debe seleccionar un arriendo")
            return

        dialog = DialogAbonosArriendo(arriendo_id, self)
        dialog.exec()

    def abrir_gastos_arriendo(self):
        arriendo_id = self.obtener_arriendo_seleccionado()
        if not arriendo_id:
            QMessageBox.warning(self, "Error", "Debe seleccionar un arriendo")
            return
        dialog = DialogGastosArriendo(arriendo_id, self)
        dialog.exec()

    def generar_excel_finanzas(self, datos, ruta, mes_desde=1, mes_hasta=12, factor_actualizacion=None):
        # -----------------------------
        # Crear workbook
        # -----------------------------
        wb = Workbook()
        wb.remove(wb.active)
        peso = '$#,##0'

        df = pd.DataFrame(datos)

        if "id" in df.columns:
            df.drop(columns=["id"], inplace=True)

        print("Número total de registros:", len(df))
        print("Datos completos:")
        for i, row in df.iterrows():
            print(f"  {i}: inicio={row.get('fecha_inicio')}, termino={row.get('fecha_termino')}, estado={row.get('estado')}")

        # -----------------------------
        # Convertir fechas de forma segura
        # -----------------------------
        df["fecha_inicio"] = pd.to_datetime(df["fecha_inicio"], errors='coerce')
        df["fecha_termino"] = pd.to_datetime(df["fecha_termino"], errors='coerce')
        
        # Renombrar columnas
        df["FECHA INICIO"] = df["fecha_inicio"]
        df["FECHA TERMINO"] = df["fecha_termino"]
        
        # Crear una versión segura para filtrado
        año_actual = pd.Timestamp.now().year
        df["FECHA_TERMINO_FILTRADO"] = df["FECHA TERMINO"].fillna(pd.Timestamp(f'{año_actual}-12-31'))

        # -----------------------------
        # Renombrar columnas
        # -----------------------------
        column_rename = {
            "rol": "ROL",
            "comuna": "COMUNA",
            "direccion": "DIRECCION",
            "nombre_arrendatario": "NOMBRE ARRENDATARIO",
            "rut_arrendatario": "RUT ARRENDATARIO",
            "renta_mensual": "MONTO RENTA",
            "estado": "ESTADO",
            "nombre_propietario": "NOMBRE PROPIETARIO",
            "rut_propietario": "RUT PROPIETARIO",
            "correo_arrendatario": "CORREO ARRENDATARIO",
            "telefono_arrendatario": "TELEFONO",
            "forma_pago": "TIPO PAGO",
            "garantia": "GARANTIA",
            "amoblado": "AMOBLADO",
            "destino": "DESTINO",
            'dfl2': "DFL12",
            "naturaleza_bien_raiz": "NATURALEZA BIEN RAIZ"
        }
        
        existing_columns = {k: v for k, v in column_rename.items() if k in df.columns}
        df.rename(columns=existing_columns, inplace=True)

        # -----------------------------
        # Convertir gastos comunes
        # -----------------------------
        if "gastos_comunes" in df.columns:
            df["GASTO COMUN"] = df["gastos_comunes"]

        # -----------------------------
        # Obtener todos los años únicos de las fechas
        # -----------------------------
        años = set()
        
        if not df.empty and "FECHA INICIO" in df.columns:
            años_inicio = df["FECHA INICIO"].dt.year.dropna().unique()
            años.update(años_inicio.astype(int))
        
        if not df.empty and "FECHA TERMINO" in df.columns:
            años_termino = df["FECHA TERMINO"].dt.year.dropna().unique()
            años.update(años_termino.astype(int))
        
        # Agregar años intermedios para contratos de larga duración
        if not df.empty:
            for _, row in df.iterrows():
                inicio = row["FECHA INICIO"]
                termino = row["FECHA_TERMINO_FILTRADO"]
                
                if pd.notnull(inicio) and pd.notnull(termino):
                    año_inicio = inicio.year
                    año_termino = termino.year
                    
                    # Agregar todos los años entre inicio y término
                    for año_intermedio in range(min(año_inicio, año_termino), max(año_inicio, año_termino) + 1):
                        años.add(año_intermedio)
        
        años = sorted(list(años), reverse=True)
        
        print(f"\nAños encontrados (incluyendo intermedios): {años}")
        print(f"Rango de meses seleccionado: {mes_desde} a {mes_hasta}")

        # -----------------------------
        # Estilos
        # -----------------------------
        bold = Font(bold=True)
        header_fill = PatternFill(start_color="D9D9D9", fill_type="solid")
        center = Alignment(horizontal="center", vertical="center")
        border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin")
        )
        fill_verde = PatternFill(start_color="C6EFCE", fill_type="solid")
        fill_rojo = PatternFill(start_color="FFC7CE", fill_type="solid")

        hojas_por_año = {}

        # ==============================
        # Generar hojas por año
        # ==============================
        for año in años:
            ws = wb.create_sheet(f"Arriendos {año}")
            row_cursor = 3

            # Título
            ws.merge_cells("A1:R1")
            t = ws["A1"]
            t.value = f"Arriendos del Año {año}"
            t.font = Font(size=16, bold=True)
            t.alignment = center

            # -----------------------------
            # Función para insertar tablas
            # -----------------------------
            def insertar_tabla(df_tabla, titulo, fill_fila):
                nonlocal row_cursor
                if df_tabla.empty:
                    print(f"  {titulo}: DataFrame vacío, no se inserta")
                    return
                
                print(f"  {titulo}: Insertando {len(df_tabla)} registros")
                
                # Columnas a mostrar
                columnas_a_mostrar = [
                    'ROL', 'COMUNA', 'DIRECCION', 'NOMBRE ARRENDATARIO', 
                    'RUT ARRENDATARIO', 'MONTO RENTA', 'ESTADO', 
                    'FECHA INICIO', 'FECHA TERMINO', 'NOMBRE PROPIETARIO',
                    'RUT PROPIETARIO', 'CORREO ARRENDATARIO', 'TELEFONO',
                    'TIPO PAGO', 'GASTO COMUN', 'NATURALEZA BIEN RAIZ',
                    'GARANTIA', 'AMOBLADO', 'DFL12', 'DESTINO'
                ]
                
                columnas_existentes = [col for col in columnas_a_mostrar if col in df_tabla.columns]
                df_a_mostrar = df_tabla[columnas_existentes].copy()
                
                num_columns = len(columnas_existentes)
                
                # Título de la tabla
                ws.merge_cells(start_row=row_cursor, start_column=1,
                            end_row=row_cursor, end_column=num_columns)
                cell_title = ws.cell(row=row_cursor, column=1, value=titulo)
                cell_title.font = Font(bold=True, size=14)
                cell_title.alignment = center
                row_cursor += 1
                
                # Encabezados
                for col_idx, col_name in enumerate(columnas_existentes, 1):
                    cell = ws.cell(row=row_cursor, column=col_idx, value=col_name)
                    cell.font = bold
                    cell.fill = header_fill
                    cell.alignment = center
                    cell.border = border
                
                row_cursor += 1
                
                # Datos
                for _, row_data in df_a_mostrar.iterrows():
                    for col_idx, col_name in enumerate(columnas_existentes, 1):
                        value = row_data[col_name]
                        
                        if 'FECHA' in col_name and pd.notnull(value):
                            if isinstance(value, pd.Timestamp):
                                value = value.strftime('%d-%m-%Y')
                        
                        if col_name == 'MONTO RENTA' and pd.notnull(value):
                            try:
                                value = float(value)
                            except:
                                pass
                        
                        cell = ws.cell(row=row_cursor, column=col_idx, value=value)
                        cell.fill = fill_fila
                        cell.border = border
                        
                        if col_name == 'MONTO RENTA':
                            cell.number_format = peso
                    
                    row_cursor += 1
                
                row_cursor += 2

            # -----------------------------
            # Filtrar arriendos para este año - NUEVA LÓGICA
            # -----------------------------
            df_vigentes = pd.DataFrame()
            df_finalizados = pd.DataFrame()
            df_continuidad = pd.DataFrame()
            df_activos_en_el_año = pd.DataFrame()  # NUEVO: todos los contratos activos durante el año
            
            if not df.empty:
                try:
                    print(f"\n=== FILTRANDO PARA AÑO {año} ===")
                    
                    # 1. Todos los contratos que estuvieron activos durante el año
                    # Esto incluye: contratos que comenzaron antes o durante el año Y terminaron después o durante el año
                    mask_activos = (
                        (df["FECHA INICIO"].dt.year <= año) &
                        (df["FECHA_TERMINO_FILTRADO"].dt.year >= año)
                    )
                    
                    if mask_activos.any():
                        df_temp = df[mask_activos].copy()
                        # Filtrar por mes de inicio
                        df_temp = df_temp[df_temp["FECHA INICIO"].dt.month.between(mes_desde, mes_hasta)]
                        df_activos_en_el_año = df_temp
                    
                    # 2. De los activos, separar por estado
                    if not df_activos_en_el_año.empty:
                        # Vigentes en el año (estado = "Vigente")
                        df_vigentes = df_activos_en_el_año[df_activos_en_el_año["ESTADO"] == "Vigente"].copy()
                        
                        # Finalizados en el año (estado != "Vigente" Y terminaron en este año)
                        mask_finalizados_en_año = (
                            (df_activos_en_el_año["ESTADO"] != "Vigente") &
                            (df_activos_en_el_año["FECHA TERMINO"].dt.year == año)
                        )
                        df_finalizados = df_activos_en_el_año[mask_finalizados_en_año].copy()
                        
                        # Continuidad: contratos vigentes que comenzaron antes del año
                        mask_continuidad = (
                            (df_activos_en_el_año["ESTADO"] == "Vigente") &
                            (df_activos_en_el_año["FECHA INICIO"].dt.year < año)
                        )
                        df_continuidad = df_activos_en_el_año[mask_continuidad].copy()
                    
                    print(f"Para año {año}:")
                    print(f"  Total activos en el año: {len(df_activos_en_el_año)}")
                    print(f"  Vigentes: {len(df_vigentes)}")
                    print(f"  Finalizados: {len(df_finalizados)}")
                    print(f"  Continuidad: {len(df_continuidad)}")
                    
                    # Mostrar detalles
                    if not df_activos_en_el_año.empty:
                        print(f"  Contratos activos:")
                        for i, row in df_activos_en_el_año.iterrows():
                            print(f"    - {row.get('DIRECCION', 'N/A')}: {row['FECHA INICIO'].strftime('%d-%m-%Y')} a {row['FECHA TERMINO'].strftime('%d-%m-%Y') if pd.notnull(row['FECHA TERMINO']) else 'Vigente'} ({row['ESTADO']})")
                    
                except Exception as e:
                    print(f"Error filtrando para año {año}: {e}")
                    import traceback
                    traceback.print_exc()

            # Insertar tablas - MOSTRAR TODOS LOS ACTIVOS
            if not df_activos_en_el_año.empty:
                # Separar por estado para colores diferentes
                df_activos_vigentes = df_activos_en_el_año[df_activos_en_el_año["ESTADO"] == "Vigente"].copy()
                df_activos_finalizados = df_activos_en_el_año[df_activos_en_el_año["ESTADO"] != "Vigente"].copy()
                
                # Mostrar todos los activos (con colores según estado)
                insertar_tabla(df_activos_vigentes, "ARRIENDOS VIGENTES EN EL AÑO", fill_verde)
                insertar_tabla(df_activos_finalizados, "ARRIENDOS FINALIZADOS EN EL AÑO", fill_rojo)
                
                # Mostrar continuidad si hay
                if not df_continuidad.empty:
                    insertar_tabla(df_continuidad, f"ARRIENDOS VIGENTES DESDE {año-1}", fill_verde)
            else:
                # Si no hay activos, mostrar tablas vacías
                insertar_tabla(df_vigentes, "ARRIENDOS VIGENTES", fill_verde)
                insertar_tabla(df_finalizados, "ARRIENDOS FINALIZADOS", fill_rojo)
                insertar_tabla(df_continuidad, f"ARRIENDOS VIGENTES DESDE {año-1}", fill_verde)

            hojas_por_año[año] = (ws, row_cursor)

        # ==============================
        # Factor de actualización
        # ==============================
       # ==============================

        for año in años:
            ws, row_cursor = hojas_por_año[año]
            if factor_actualizacion and not df.empty:
                try:
                    # Filtrar SOLO contratos VIGENTES para este año específico
                    df_vigentes_año = df[
                        (df["ESTADO"] == "Vigente") &  # SOLO VIGENTES
                        (df["FECHA INICIO"].dt.year <= año) &
                        (df["FECHA_TERMINO_FILTRADO"].dt.year >= año) &
                        (df["FECHA INICIO"].dt.month.between(mes_desde, mes_hasta))
                    ]
                    
                    if not df_vigentes_año.empty:
                        ws.merge_cells(start_row=row_cursor, start_column=1,
                                    end_row=row_cursor, end_column=5)
                        ws.cell(row=row_cursor, column=1, 
                            value=f"DETALLE ARRIENDOS VIGENTES CON FACTOR DE ACTUALIZACIÓN - AÑO {año}")
                        ws.cell(row=row_cursor, column=1).font = Font(bold=True)
                        row_cursor += 1

                        headers = ["Mes","Propiedad","Renta mensual","Factor","Renta actualizada"]
                        for col, h in enumerate(headers, 1):
                            c = ws.cell(row=row_cursor, column=col, value=h)
                            c.font = bold
                            c.fill = header_fill
                            c.alignment = center
                            c.border = border
                        row_cursor += 1

                        total_general_actualizado = 0
                        for f in factor_actualizacion:
                            mes = f["mes"]
                            factor = f["factor"]
                            for _, r in df_vigentes_año.iterrows():
                                renta = r.get("MONTO RENTA", 0)
                                propiedad = r.get("DIRECCION", "")
                                total_actualizado = renta * factor
                                total_general_actualizado += total_actualizado
                                ws.cell(row=row_cursor, column=1, value=mes)
                                ws.cell(row=row_cursor, column=2, value=propiedad)
                                ws.cell(row=row_cursor, column=3, value=int(renta))
                                ws.cell(row=row_cursor, column=4, value=factor)
                                ws.cell(row=row_cursor, column=5, value=int(total_actualizado))
                                ws.cell(row=row_cursor, column=3).number_format = peso
                                ws.cell(row=row_cursor, column=5).number_format = peso
                                for col in range(1, 6):
                                    ws.cell(row=row_cursor, column=col).border = border
                                row_cursor += 1

                        ws.merge_cells(start_row=row_cursor, start_column=1, end_row=row_cursor, end_column=4)
                        c = ws.cell(row=row_cursor, column=1, value="TOTAL RENTAS VIGENTES ACTUALIZADAS")
                        c.font = Font(bold=True)
                        c.alignment = center
                        c.border = border
                        total_cell = ws.cell(row=row_cursor, column=5, value=int(total_general_actualizado))
                        total_cell.number_format = peso
                        row_cursor += 2
                        
                        # También mostrar resumen
                        ws.merge_cells(start_row=row_cursor, start_column=1, end_row=row_cursor, end_column=5)
                        ws.cell(row=row_cursor, column=1, 
                            value=f"Resumen: {len(df_vigentes_año)} arriendo(s) vigente(s) actualizado(s)")
                        ws.cell(row=row_cursor, column=1).font = Font(italic=True)
                        row_cursor += 2
                        
                except Exception as e:
                    print(f"Error agregando factor de actualización para año {año}: {e}")

        # -----------------------------
        # Ajustar anchos de columnas
        # -----------------------------
        for ws in wb.worksheets:
            for col_idx, column_cells in enumerate(ws.columns, 1):
                try:
                    values = [str(cell.value) for cell in column_cells if cell.value is not None]
                    if values:
                        length = max(len(v) for v in values) + 3
                        ws.column_dimensions[get_column_letter(col_idx)].width = min(length, 50)
                    else:
                        ws.column_dimensions[get_column_letter(col_idx)].width = 15
                except Exception:
                    ws.column_dimensions[get_column_letter(col_idx)].width = 15

        try:
            wb.save(ruta)
            print(f"\nExcel guardado exitosamente en: {ruta}")
            return True
        except Exception as e:
            print(f"Error guardando el archivo: {e}")
            return False


    def exportar_excel(self):
        # -----------------------------
        # Obtener datos de Supabase (lista)
        # -----------------------------
        datos = obtener_arriendos_finanzas()  # lista de diccionarios

        if not datos:
            QMessageBox.warning(self, "Sin datos", "No se encontraron arriendos para exportar.")
            return

        # -----------------------------
        # Convertir a DataFrame
        # -----------------------------
        df = pd.DataFrame(datos)

        # -----------------------------
        # Convertir fechas de forma segura
        # -----------------------------
        df["FECHA INICIO"] = pd.to_datetime(df.get("fecha_inicio"), errors='coerce')
        df["FECHA TERMINO"] = pd.to_datetime(df.get("fecha_termino"), errors='coerce')
        df["FECHA_TERMINO_FILTRADO"] = df["FECHA TERMINO"].fillna(pd.Timestamp.max)

        # -----------------------------
        # Mostrar filas con fechas inválidas (depuración)
        # -----------------------------
        for i, row in df.iterrows():
            if pd.isna(row["FECHA INICIO"]):
                print(f"Fila {i} FECHA_INICIO inválida: {row.get('fecha_inicio')}")
            if pd.isna(row["FECHA TERMINO"]):
                print(f"Fila {i} FECHA_TERMINO inválida o nula: {row.get('fecha_termino')}")

        # -----------------------------
        # Selección de mes desde/hasta
        # -----------------------------
        mes_desde = self.cmb_mes_desde.currentData()
        mes_hasta = self.cmb_mes_hasta.currentData()

        print(f"Meses seleccionados: desde {mes_desde} hasta {mes_hasta}")

        if mes_desde > mes_hasta:
            QMessageBox.warning(self, "Error", "El mes 'desde' no puede ser mayor al mes 'hasta'")
            return

        # -----------------------------
        # Diálogo para guardar archivo
        # -----------------------------
        ruta, _ = QFileDialog.getSaveFileName(
            self,
            "Guardar Excel",
            "consolidado_arriendos",
            "Archivos Excel (*.xlsx)"
        )
        if not ruta:
            return
        if not ruta.endswith(".xlsx"):
            ruta += ".xlsx"

 
        try:
            print("=== INICIANDO GENERACIÓN DE EXCEL ===")
            self.generar_excel_finanzas(
                df,
                ruta,
                factor_actualizacion=getattr(self, "factores_actualizacion", None),
                mes_desde=mes_desde,
                mes_hasta=mes_hasta
            )
            QMessageBox.information(self, "Éxito", f"Excel generado correctamente:\n{ruta}")
        except Exception as e:
            error_msg = f"Error al generar Excel:\n{str(e)}"
            print(error_msg)
            import traceback
            traceback.print_exc()
            QMessageBox.critical(self, "Error", error_msg)


    def abrir_formulario_arriendo(self):
        self.formulario_arriendo = FormularioArriendo()
        self.formulario_arriendo.arriendo_guardado.connect(self.cargar_arriendos)
        self.formulario_arriendo.show()

    def abrir_factor_actualizacion(self):
        dialog = DialogFactorActualizacion(self)
        if dialog.exec():
            self.factores_actualizacion = dialog.resultado

                

class DoubleClickButton(QPushButton):
    doubleClicked = Signal()

    def mouseDoubleClickEvent(self, event):
        self.doubleClicked.emit()


        
    

class FormularioArriendo(QWidget):
    arriendo_guardado = Signal()

    def __init__(self,arriendo_id=None, parent=None):
        super().__init__(parent)
        self.arriendo_id = arriendo_id
        self.rentas_tributables = []

        self.setWindowTitle("Editar Arriendo" if arriendo_id else "Nuevo Arriendo")
        self.resize(900, 700)
        self.detalle_arriendo = {}
        self.evaluacion_meses = {}

        # Layout principal con scroll
        layout_principal = QVBoxLayout(self)
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        contenido = QWidget()
        scroll.setWidget(contenido)
        layout_principal.addWidget(scroll)
        
        # Layout del formulario
        layout_form = QVBoxLayout(contenido)
        
        # Pestañas para organizar el formulario
        self.tabs = QTabWidget()
        layout_form.addWidget(self.tabs)


        tab_arrendador = QWidget()
        self.tabs.addTab(tab_arrendador, "Arrendador")
        self.setup_tab_arrendador(tab_arrendador)

        tab_arrendatario = QWidget()
        self.tabs.addTab(tab_arrendatario , "Arrendatario")
        self.setup_tab_arrendatario(tab_arrendatario )

        tab_propiedad = QWidget()
        self.tabs.addTab(tab_propiedad, "Propiedad")
        self.setup_tab_propiedad(tab_propiedad)

        tab_evaluacion = QWidget()
        self.tabs.addTab(tab_evaluacion, "Evaluación")
        self.setup_tab_evaluacion(tab_evaluacion)
        self.setup_tab_evaluacion_index = self.tabs.indexOf(tab_evaluacion)

        tab_evaluacion_independiente = QWidget()
        self.tabs.addTab(tab_evaluacion_independiente, "Evaluación")
        self.setup_tab_evaluacion_independiente(tab_evaluacion_independiente)
        self.setup_tab_evaluacion_independiente_index = self.tabs.indexOf(tab_evaluacion_independiente )

        tab_arriendo = QWidget()
        self.tabs.addTab(tab_arriendo, "Arriendo")
        self.setup_tab_arriendo(tab_arriendo)


        self.cmb_tipo_trabajador.currentTextChanged.connect(self.toggle_evaluacion_tab)
        self.toggle_evaluacion_tab(self.cmb_tipo_trabajador.currentText())

        self.cmb_tipo_trabajador.currentTextChanged.connect(self.toggle_evaluacion_independiente_tab)
        self.toggle_evaluacion_independiente_tab(self.cmb_tipo_trabajador.currentText())

        


        btn_guardar = DoubleClickButton("Guardar Arriendo")
        btn_guardar.doubleClicked.connect(self.guardar_arriendo)
        layout_form.addWidget(btn_guardar)
        
        # Estilo para campos obligatorios
        self.setStyleSheet("""
            QLabel[obligatorio="true"] {
                font-weight: bold;
            }
        """)

        if self.arriendo_id:
            self.detalle_arriendo = obtener_detalle_arriendo(self.arriendo_id)
            if not self.detalle_arriendo:
                layout_principal.addWidget(QLabel("No se encontraron detalles para este arriendo"))
            else:
                self.cargar_datos_arriendo(self.arriendo_id)

    def money(self, value):
            return f"{int(safe_numeric(value)):,}".replace(",", ".")

    def percent(self, value):
            return f"{safe_numeric(value) * 100:.1f}%"

    def safe_text(self, value):
            return "" if value in (None, "") else str(value)

            


    def setup_tab_arrendador(self, tab):
        layout = QFormLayout(tab)
        
        # Campos del comprador
        self.txt_arren_nombre = QLineEdit()
        self.txt_arren_rut = QLineEdit()
        self.txt_arren_direccion = QLineEdit()
        self.txt_arren_telefono = QLineEdit()
        self.txt_arren_email = QLineEdit()
     
        
        # Agregar campos

        layout.addRow(self.crear_label("Nombre:", True), self.txt_arren_nombre)
        layout.addRow(self.crear_label("RUT:", True), self.txt_arren_rut)
        layout.addRow(self.crear_label("Dirección:"), self.txt_arren_direccion)
        layout.addRow(self.crear_label("Teléfono:"), self.txt_arren_telefono)
        layout.addRow(self.crear_label("Email:"), self.txt_arren_email)



    def setup_tab_arrendatario(self, tab):
        layout = QFormLayout(tab)
        
        # Campos del arrendatario
        self.txt_arrendatario_nombre = QLineEdit()
        self.txt_arrendatario_rut = QLineEdit()
        self.txt_arrendatario_direccion = QLineEdit()
        self.txt_arrendatario_telefono = QLineEdit()
        self.txt_arrendatario_email = QLineEdit()
        self.cmb_tipo_trabajador = QComboBox()
        self.cmb_tipo_trabajador.addItems(["Dependiente", "Independiente", "Otro"])
        self.txt_antiguedad_laboral = QLineEdit()
        self.cmb_dicom = QComboBox()
        self.cmb_dicom.addItems(["Sí","No"])
        self.txt_comentarios = QTextEdit()
        self.cmb_evaluacion_estado = QComboBox()
        self.cmb_evaluacion_estado.addItems(["Pendiente","Aprobado","Rechazado"])

        self.fecha_evaluacion = QDateEdit(QDate.currentDate())
        self.fecha_evaluacion.setCalendarPopup(True)

        self.txt_renta = QLineEdit()

     
        
        # Agregar campos

        layout.addRow(self.crear_label("Nombre:", True), self.txt_arrendatario_nombre)
        layout.addRow(self.crear_label("RUT:", True), self.txt_arrendatario_rut)
        layout.addRow(self.crear_label("Dirección:"), self.txt_arrendatario_direccion)
        layout.addRow(self.crear_label("Teléfono:"), self.txt_arrendatario_telefono)
        layout.addRow(self.crear_label("Email:"), self.txt_arrendatario_email)
        
        layout.addRow(self.crear_label("Tipo de Trabajador:"), self.cmb_tipo_trabajador)
        layout.addRow(self.crear_label("Antigüedad Laboral(Meses):"), self.txt_antiguedad_laboral)

        # --- Evaluación ---
        layout.addRow(self.crear_label("¿Tiene DICOM?"), self.cmb_dicom)
        layout.addRow(self.crear_label("Comentarios:"), self.txt_comentarios)
        layout.addRow(self.crear_label("Estado de Evaluación:"), self.cmb_evaluacion_estado)
        layout.addRow(self.crear_label("Fecha Evaluación:"), self.fecha_evaluacion)
        layout.addRow(self.crear_label("Renta Mensual:"), self.txt_renta)



    def toggle_evaluacion_tab(self, tipo_trabajador):
        if not hasattr(self,'setup_tab_evaluacion_index'):
            return
        mostrar = tipo_trabajador == "Dependiente"
        self.tabs.setTabVisible(self.setup_tab_evaluacion_index, mostrar)

    def toggle_evaluacion_independiente_tab(self, tipo_trabajador):
        if not hasattr(self, 'setup_tab_evaluacion_independiente_index'):
            return
        mostrar = tipo_trabajador == "Independiente"
        self.tabs.setTabVisible(self.setup_tab_evaluacion_independiente_index, mostrar)
                

    def setup_tab_propiedad(self, tab):
        layout = QFormLayout(tab)
        
        # Campos de propiedad
        self.txt_prop_codigo = QLineEdit()
        self.txt_prop_direccion = QLineEdit()
        self.txt_prop_rol = QLineEdit()
        self.txt_prop_comuna = QLineEdit()
        
        # Documentación
        
        self.cmb_dominio_vigente = QComboBox()
        self.cmb_dominio_vigente.addItems(["Si Posee Documento", "No Posee Documento"])


        # Agregar campos
        layout.addRow(self.crear_label("Código interno:", True), self.txt_prop_codigo)
        layout.addRow(self.crear_label("Dirección:", True), self.txt_prop_direccion)
        layout.addRow(self.crear_label("ROL:", True), self.txt_prop_rol)
        layout.addRow(self.crear_label("Comuna:", True), self.txt_prop_comuna)
        
        # Documentación
        layout.addRow(QLabel("<b>Documentación:</b>"))
        layout.addRow(self.crear_label("Dominio vigente:"), self.cmb_dominio_vigente)


    def setup_tab_evaluacion(self, tab):
        layout = QFormLayout(tab)

        # Contenedor para centrar la tabla
        contenedor = QWidget()
        contenedor_layout = QVBoxLayout(contenedor)
        contenedor_layout.setAlignment(Qt.AlignCenter)

        # Crear tabla

        self.fecha_evaluacion_arrendatario = QDateEdit(QDate.currentDate())
        self.fecha_evaluacion_arrendatario.setCalendarPopup(True)

        layout.addRow(self.crear_label("Fecha Evaluación:"), self.fecha_evaluacion_arrendatario)

        self.txt_impuesto_renta = QLineEdit()
        layout.addRow(self.crear_label("Impuesto Renta"), self.txt_impuesto_renta)

        self.btn_calcular_impuesto = DoubleClickButton("Calcular Impuesto Renta")
        self.btn_calcular_impuesto.doubleClicked.connect(self.abrir_dialogo_impuesto)
        layout.addRow(self.btn_calcular_impuesto)

        self.tbl_evaluacion = QTableWidget()
        self.tbl_evaluacion.setRowCount(10)
        self.tbl_evaluacion.setColumnCount(5)

        # Tamaño mínimo más grande
        self.tbl_evaluacion.setMinimumWidth(700)
        self.tbl_evaluacion.setMinimumHeight(300)

        # Encabezados
        self.tbl_evaluacion.setHorizontalHeaderLabels([
            "Concepto", "Mes 1", "Mes 2", "Mes 3", "Resultado"
        ])

        # Centrar texto de encabezados
        header = self.tbl_evaluacion.horizontalHeader()
        header.setDefaultAlignment(Qt.AlignCenter)

        conceptos = [
            "Sueldo Base (SB)",
            "Gratificación (GT)",
            "TH Imponibles (THIMP)",
            "Locomoción / Movilización (LOCMOV)",
            "TH No Imponibles (THNI)",
            "Total Haberes",
            "Descuentos Legales",
            "Descuentos Varios",
            "Anticipos",
            "Líquido a Pago"
        ]

        # Llenar columna de conceptos
        for i, concepto in enumerate(conceptos):
            item = QTableWidgetItem(concepto)
            item.setFlags(item.flags() ^ Qt.ItemIsEditable)
            item.setTextAlignment(Qt.AlignVCenter | Qt.AlignLeft)
            self.tbl_evaluacion.setItem(i, 0, item)

        # Ajustar automáticamente tamaños de columnas
        self.tbl_evaluacion.resizeColumnsToContents()
        header = self.tbl_evaluacion.horizontalHeader()

        header.setSectionResizeMode(0, QHeaderView.Stretch)
        header.setSectionResizeMode(1, QHeaderView.Stretch)
        header.setSectionResizeMode(2, QHeaderView.Stretch)
        header.setSectionResizeMode(3, QHeaderView.Stretch)
        header.setSectionResizeMode(4, QHeaderView.Stretch)
        header.setSectionResizeMode(5, QHeaderView.Stretch)
        

        # Añadir tabla al contenedor centrado
        contenedor_layout.addWidget(self.tbl_evaluacion)

        botones_layout = QHBoxLayout()
        botones_layout.setAlignment(Qt.AlignCenter)

        self.btn_mes1 = DoubleClickButton("Ingresar Mes 1")
        self.btn_mes2 = DoubleClickButton("Ingresar Mes 2")
        self.btn_mes3 = DoubleClickButton("Ingresar Mes 3")

        self.btn_mes1.doubleClicked.connect(lambda: self.abrir_formulario_mes(1))
        self.btn_mes2.doubleClicked.connect(lambda: self.abrir_formulario_mes(2))
        self.btn_mes3.doubleClicked.connect(lambda: self.abrir_formulario_mes(3))

        botones_layout.addWidget(self.btn_mes1)
        botones_layout.addWidget(self.btn_mes2)
        botones_layout.addWidget(self.btn_mes3)

        contenedor_layout.addLayout(botones_layout)

        # Agregar contenedor al layout final
        layout.addRow(contenedor)


    def setup_tab_evaluacion_independiente(self, tab):
        layout = QVBoxLayout(tab)

        # --- Tabla mensual ---
        self.tbl_independiente = QTableWidget()
        self.tbl_independiente.setRowCount(12)
        self.tbl_independiente.setColumnCount(6)

        headers = [
            "Mes",
            "IVA Débito (538)",
            "IVA Crédito (537)",
            "Ventas Netas Est.",
            "Compras Netas Est.",
            "Excedente Mes"
        ]
        self.tbl_independiente.setHorizontalHeaderLabels(headers)

        # 🔹 Headers centrados y en negrita
        header = self.tbl_independiente.horizontalHeader()
        font = QFont()
        font.setBold(True)
        header.setFont(font)

        for i in range(self.tbl_independiente.columnCount()):
            header.setSectionResizeMode(i, QHeaderView.Stretch)

        header.setDefaultAlignment(Qt.AlignCenter)

        # 🔹 Estética general
        self.tbl_independiente.verticalHeader().setVisible(False)
        self.tbl_independiente.setAlternatingRowColors(True)
        self.tbl_independiente.setEditTriggers(QTableWidget.DoubleClicked | QTableWidget.EditKeyPressed)

        self.tbl_independiente.setSelectionBehavior(QTableWidget.SelectRows)

        # 🔹 Centrar TODAS las celdas
        for row in range(12):
            for col in range(6):
                item = QTableWidgetItem("")
                item.setTextAlignment(Qt.AlignCenter)
                self.tbl_independiente.setItem(row, col, item)

        layout.addWidget(self.tbl_independiente)

                # --- Acciones ---
        acciones_frame = QFrame()
        acciones_frame.setFrameShape(QFrame.StyledPanel)
     
        acciones_layout = QHBoxLayout(acciones_frame)
        acciones_layout.setAlignment(Qt.AlignRight)

        btn_excel = DoubleClickButton("📄 Generar Excel")
        btn_excel.setFixedHeight(36)
        btn_excel.setCursor(Qt.PointingHandCursor)

        btn_excel.doubleClicked.connect(self.obtener_ruta_excel)

        acciones_layout.addWidget(btn_excel)
        layout.addWidget(acciones_frame)



        # --- Resumen ---
        resumen_layout = QHBoxLayout()

        self.lbl_ventas = QLabel("Ventas Anuales: $0")
        self.lbl_compras = QLabel("Compras Anuales: $0")
        self.lbl_excedente = QLabel("Excedente Anual: $0")
        self.lbl_renta = QLabel("Renta Mensual Estimada: $0")

        for lbl in (
            self.lbl_ventas,
            self.lbl_compras,
            self.lbl_excedente,
            self.lbl_renta
        ):
            lbl.setAlignment(Qt.AlignCenter)
            lbl.setStyleSheet("font-weight: bold;")
            resumen_layout.addWidget(lbl)

        layout.addLayout(resumen_layout)

        self.configurar_edicion_tabla_independiente()

        self.tbl_independiente.itemChanged.connect(
                self.recalcular_evaluacion_independiente
            )
        
        self.cargar_meses_independiente()


    
    def configurar_edicion_tabla_independiente(self):
        for row in range(12):
            for col in range(6):
                item = self.tbl_independiente.item(row, col)
                flags = item.flags()

                if col in (1, 2):  # IVA Débito y Crédito
                    item.setFlags(flags | Qt.ItemIsEditable)
                else:
                    item.setFlags(flags & ~Qt.ItemIsEditable)


    def recalcular_evaluacion_independiente(self, item):
        if item.column() not in (1, 2):
            return

        meses = []

        for row in range(12):
            try:
                iva_debito = float(
                    self.tbl_independiente.item(row, 1).text()
                    .replace(".", "").replace(",", ".") or 0
                )
                iva_credito = float(
                    self.tbl_independiente.item(row, 2).text()
                    .replace(".", "").replace(",", ".") or 0
                )
            except ValueError:
                iva_debito = 0
                iva_credito = 0

            meses.append({
                "iva_debito": iva_debito,
                "iva_credito": iva_credito
            })

        resultado = calcular_evaluacion_independiente(meses)

        self.actualizar_tabla_independiente(resultado)
        self.actualizar_resumen_independiente(resultado)

    def actualizar_tabla_independiente(self, resultado):
        self.tbl_independiente.blockSignals(True)

        for row, detalle in enumerate(resultado["detalle"]):
            self.tbl_independiente.item(row, 3).setText(
                f"{detalle['ventas']:,.0f}"
            )
            self.tbl_independiente.item(row, 4).setText(
                f"{detalle['compras']:,.0f}"
            )
            self.tbl_independiente.item(row, 5).setText(
                f"{detalle['excedente']:,.0f}"
            )

        self.tbl_independiente.blockSignals(False)

    def cargar_meses_independiente(self):
        MESES = [
            "Enero", "Febrero", "Marzo", "Abril",
            "Mayo", "Junio", "Julio", "Agosto",
            "Septiembre", "Octubre", "Noviembre", "Diciembre"
        ]

        for row, mes in enumerate(MESES):
            item_mes = QTableWidgetItem(mes)
            item_mes.setFlags(item_mes.flags() & ~Qt.ItemIsEditable)
            item_mes.setTextAlignment(Qt.AlignCenter)
            self.tbl_independiente.setItem(row, 0, item_mes)



    def actualizar_resumen_independiente(self, resultado):
        self.lbl_ventas.setText(
            f"Ventas Anuales: ${resultado['ventas_anuales']:,.0f}"
        )
        self.lbl_compras.setText(
            f"Compras Anuales: ${resultado['compras_anuales']:,.0f}"
        )
        self.lbl_excedente.setText(
            f"Excedente Anual: ${resultado['excedente_anual']:,.0f}"
        )
        self.lbl_renta.setText(
            f"Renta Mensual Estimada: ${resultado['renta_mensual']:,.0f}"
        )


    def abrir_formulario_mes(self, mes):
        dlg = FormularioMes(mes, self, rut_arrendatario=self.txt_arrendatario_rut.text())

        if dlg.exec():
            resultado = dlg.resultado_final   # ← ya trae todo

            datos = resultado["datos"]
            calculos = resultado["calculos"]

            # Guardas lo que quieras
            self.evaluacion_meses[mes] = datos

            # Rellenas la tabla usando los cálculos
            resultados_completos = {**datos, **calculos}

            self.actualizar_tabla_evaluacion(mes, resultados_completos)

    def calcular_resultados_evaluacion(self):
        """
        Recorre la tabla de evaluación y calcula la columna Resultado.
        Si los valores de Mes1, Mes2 y Mes3 son iguales → usa ese valor.
        Si hay cambios → calcula el promedio.
        """
        COL_M1 = 1
        COL_M2 = 2
        COL_M3 = 3
        COL_RES = 4

        for fila in range(self.tbl_evaluacion.rowCount()):

            valores = []

            # Obtener los valores de cada mes
            for col in (COL_M1, COL_M2, COL_M3):
                item = self.tbl_evaluacion.item(fila, col)

                if item and item.text().strip():
                    texto = item.text().replace(".", "").replace(",", "")
                    try:
                        valor = float(texto)
                    except:
                        valor = 0
                else:
                    valor = 0

                valores.append(valor)

            v1, v2, v3 = valores

            # --- LÓGICA DE PROMEDIO O IGUALDAD ---
            if v1 == v2 == v3:
                resultado = v1
            else:
                resultado = sum(valores) / 3

            # Crear item resultado
            item_res = QTableWidgetItem(f"{resultado:,.0f}")
            item_res.setTextAlignment(Qt.AlignCenter)
            item_res.setFlags(item_res.flags() ^ Qt.ItemIsEditable)

            self.tbl_evaluacion.setItem(fila, COL_RES, item_res)


    def actualizar_tabla_evaluacion(self, mes, resultados):
        col = mes  

        # print("CALCULOS MES:", resultados)

        mapping = [
            ("sueldo_base", 0),
            ("gratificacion", 1),
            ("th_imponibles", 2),
            ("locomocion", 3),
            ("th_no_imponibles", 4),
            ("total_haberes",5),
            ("descuentos_legales", 6),
            ("descuentos_varios", 7),
            ("anticipo", 8),
            ("liquido", 9)
        ]

        for clave, fila in mapping:
            valor = resultados.get(clave, 0)
            item = QTableWidgetItem(f"{valor:,.0f}")
            item.setTextAlignment(Qt.AlignCenter)
            self.tbl_evaluacion.setItem(fila, col, item)

        self.calcular_resultados_evaluacion()

    
    def obtener_diccionario_evaluacion(self):
        col = 4  # columna Resultado

        mapping = [
            "sueldo_base",
            "gratificacion",
            "total_imponible",
            "locomocion",
            "total_no_imponible",
            "total_haberes",
            "descuentos_legales",
            "desc_varios",
            "anticipo",
            "liquido_pago",
        ]

        datos = {}

        for fila, clave in enumerate(mapping):
            item = self.tbl_evaluacion.item(fila, col)

            if item and item.text().strip():
                texto = item.text().replace(".", "").replace(",", "").strip()
                valor = float(texto) if texto else 0.0
            else:
                valor = 0.0

            datos[clave] = valor

        return datos
    

    def abrir_dialogo_impuesto(self):
        datos_eval = self.obtener_diccionario_evaluacion()
        renta_tributable = datos_eval.get("total_imponible", 0.0)


        dialog = ImpuestoRentaDialog(renta_tributable, self)

        if dialog.exec():
            self.txt_impuesto_renta.setText(str(dialog.resultado))


   
        

    def setup_tab_arriendo(self, tab):

        # =========================
        # Widgets
        # =========================
        self.fecha_inicio = QDateEdit(QDate.currentDate())
        self.fecha_inicio.setCalendarPopup(True)

        self.fecha_termino = QDateEdit(QDate.currentDate())
        self.fecha_termino.setCalendarPopup(True)

        self.txt_renta_mensual = QLineEdit()
        self.txt_garantia = QLineEdit()

        self.chk_gastos_comunes = QCheckBox("¿Gastos comunes incluidos?")
        self.chk_gastos_comunes.setChecked(False)

        self.cmb_estado_arriendo = QComboBox()
        self.cmb_estado_arriendo.addItems(["Vigente", "Finalizado", "Rescindido"])

        self.cmb_tipo_contrato = QComboBox()
        self.cmb_tipo_contrato.addItems(["Plazo Fijo", "Plazo Indefinido"])

        self.txt_forma_pago = QLineEdit()

        self.cmb_periodo_pago = QComboBox()
        self.cmb_periodo_pago.addItems(["Mensual", "Trimestral", "Anual"])

        self.cmb_cuenta_fm = QComboBox()
        self.cmb_cuenta_fm.addItems(["1", "2", "3", "4"])

        self.cmb_tipo_documento = QComboBox()
        self.cmb_tipo_documento.addItems(["Boleta", "Factura"])

        self.txt_nro_cuenta = QLineEdit()
        self.txt_banco_destino = QLineEdit()
        self.txt_tipo_cuenta = QLineEdit()
        self.txt_rut_para_deposito = QLineEdit()
        self.txt_titular_deposito = QLineEdit()
        self.txt_quien_deposita = QLineEdit()
        self.txt_correo_deposito = QLineEdit()

        self.txt_aseo_municipal = QLineEdit()
        self.txt_reajuste = QLineEdit()
        self.txt_ggcc = QLineEdit()
        self.txt_cuenta_ggcc = QLineEdit()

        self.txt_honorarios_porcentaje = QLineEdit()
        self.txt_honorarios_monto = QLineEdit()

        self.txt_dia_pago = QLineEdit()
        self.txt_ultimo_mes_pago = QLineEdit()
        self.txt_direccion_consulta = QLineEdit()
        self.txt_periodo_anterior = QLineEdit()
        self.txt_monto_anterior = QLineEdit()
        self.txt_naturaleza_bien_raiz = QLineEdit()

        self.cmb_dfl12 = QComboBox()
        self.cmb_dfl12.addItems(["Si", "No"])

        self.txt_destino = QLineEdit()

        self.cmb_amoblado = QComboBox()
        self.cmb_amoblado.addItems(["Si", "No"])

        self.txt_observaciones = QTextEdit()

        # =========================
        # Layout principal del tab
        # =========================
        main_layout = QVBoxLayout(tab)
        main_layout.setContentsMargins(0, 0, 0, 0)
        main_layout.setSpacing(0)

        # =========================
        # Content widget (va dentro del scroll)
        # =========================
        content_widget = QWidget()
        content_vbox = QVBoxLayout(content_widget)
        content_vbox.setContentsMargins(8, 8, 8, 8)
        content_vbox.setSpacing(10)
        content_vbox.setAlignment(Qt.AlignTop)

        # -------------------------
        # Box 1: Básicos
        # -------------------------
        box_basicos = QGroupBox("Datos básicos del arriendo")
        layout_basicos = QFormLayout()
        layout_basicos.addRow(self.crear_label("Fecha Inicio:", True), self.fecha_inicio)

        self.lbl_fecha_termino = self.crear_label("Fecha de Término:")
        layout_basicos.addRow(self.lbl_fecha_termino, self.fecha_termino)

        layout_basicos.addRow(self.crear_label("Tipo de Documento"), self.cmb_tipo_documento)
        layout_basicos.addRow(self.crear_label("Renta Mensual:"), self.txt_renta_mensual)
        layout_basicos.addRow(self.crear_label("Garantía:"), self.txt_garantia)
        layout_basicos.addRow(self.chk_gastos_comunes)
        layout_basicos.addRow(self.crear_label("Estado del arriendo:"), self.cmb_estado_arriendo)
        layout_basicos.addRow(self.crear_label("Tipo de contrato:"), self.cmb_tipo_contrato)
        box_basicos.setLayout(layout_basicos)
        box_basicos.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Maximum)
        content_vbox.addWidget(box_basicos)

        self.cmb_estado_arriendo.currentTextChanged.connect(
            self.actualizar_visibilidad_fecha_termino
        )

        # -------------------------
        # Box 2: Pagos
        # -------------------------
        box_pagos = QGroupBox("Datos de Pago")
        layout_pagos = QFormLayout()
        layout_pagos.addRow(self.crear_label("Forma de pago:"), self.txt_forma_pago)
        layout_pagos.addRow(self.crear_label("Periodo de pago:"), self.cmb_periodo_pago)
        layout_pagos.addRow(self.crear_label("Día de pago:"), self.txt_dia_pago)
        layout_pagos.addRow(self.crear_label("Cuenta FM:"), self.cmb_cuenta_fm)
        box_pagos.setLayout(layout_pagos)
        box_pagos.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Maximum)
        content_vbox.addWidget(box_pagos)

        # -------------------------
        # Box 3: Depósitos
        # -------------------------
        box_depositos = QGroupBox("Depósitos / Cuenta destino")
        layout_depositos = QFormLayout()
        layout_depositos.addRow(self.crear_label("Número de cuenta:"), self.txt_nro_cuenta)
        layout_depositos.addRow(self.crear_label("Banco destino:"), self.txt_banco_destino)
        layout_depositos.addRow(self.crear_label("Tipo de cuenta:"), self.txt_tipo_cuenta)
        layout_depositos.addRow(self.crear_label("RUT para depósito:"), self.txt_rut_para_deposito)
        layout_depositos.addRow(self.crear_label("Titular depósito:"), self.txt_titular_deposito)
        layout_depositos.addRow(self.crear_label("Quién deposita:"), self.txt_quien_deposita)
        layout_depositos.addRow(self.crear_label("Correo depósito:"), self.txt_correo_deposito)
        box_depositos.setLayout(layout_depositos)
        box_depositos.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Maximum)
        content_vbox.addWidget(box_depositos)

        # -------------------------
        # Box 4: Honorarios
        # -------------------------
        box_honorarios = QGroupBox("Honorarios")
        layout_honorarios = QFormLayout()
        layout_honorarios.addRow(self.crear_label("Honorarios (%)"), self.txt_honorarios_porcentaje)
        layout_honorarios.addRow(self.crear_label("Honorarios (monto)"), self.txt_honorarios_monto)
        box_honorarios.setLayout(layout_honorarios)
        box_honorarios.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Maximum)
        content_vbox.addWidget(box_honorarios)

        # -------------------------
        # Box 5: Información adicional
        # -------------------------
        box_extra = QGroupBox("Información adicional")
        layout_extra = QFormLayout()
        layout_extra.addRow(self.crear_label("Último mes pago:"), self.txt_ultimo_mes_pago)
        layout_extra.addRow(self.crear_label("Aseo municipal:"), self.txt_aseo_municipal)
        layout_extra.addRow(self.crear_label("Reajuste:"), self.txt_reajuste)
        layout_extra.addRow(self.crear_label("GGCC:"), self.txt_ggcc)
        layout_extra.addRow(self.crear_label("Cuenta GGCC:"), self.txt_cuenta_ggcc)
        layout_extra.addRow(self.crear_label("Dirección consulta:"), self.txt_direccion_consulta)
        layout_extra.addRow(self.crear_label("Periodo anterior:"), self.txt_periodo_anterior)
        layout_extra.addRow(self.crear_label("Monto anterior:"), self.txt_monto_anterior)
        layout_extra.addRow(self.crear_label("Naturaleza bien raíz:"), self.txt_naturaleza_bien_raiz)
        layout_extra.addRow(self.crear_label("DFL12:"), self.cmb_dfl12)
        layout_extra.addRow(self.crear_label("Destino:"), self.txt_destino)
        layout_extra.addRow(self.crear_label("Amoblado:"), self.cmb_amoblado)
        layout_extra.addRow(self.crear_label("Observaciones:"), self.txt_observaciones)
        box_extra.setLayout(layout_extra)
        box_extra.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Maximum)
        content_vbox.addWidget(box_extra)

        # Empuja todo hacia arriba
        content_vbox.addStretch()

        # =========================
        # Scroll Area
        # =========================
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setFrameStyle(0)
        scroll.setWidget(content_widget)
        scroll.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)


        self.actualizar_visibilidad_fecha_termino(
            self.cmb_estado_arriendo.currentText()
        )


        # =========================
        # Añadir al tab
        # =========================
        main_layout.addWidget(scroll, 1)



    
    def actualizar_visibilidad_fecha_termino(self, estado):
        es_vigente = estado == "Vigente"

        self.lbl_fecha_termino.setVisible(not es_vigente)
        self.fecha_termino.setVisible(not es_vigente)

        # Opcional: limpiar fecha cuando se oculta
        if es_vigente:
            self.fecha_termino.clear()

            


    def cargar_datos_arriendo(self, arriendo_id):
        try:

            def set_text(widget, value):
                widget.setText(self.safe_text(value))

            def set_money(widget, value):
                widget.setText(self.money(value))

            def set_date(widget, value):
                if value:
                    widget.setDate(QDate.fromString(value, "yyyy-MM-dd"))
                else:
                    widget.setDate(QDate.currentDate())

            # ===============================
            # ARRENDADOR
            # ===============================
            arrendador = self.detalle_arriendo.get('arrendador', {})
            if isinstance(arrendador, dict):
                set_text(self.txt_arren_nombre, arrendador.get("nombre"))
                set_text(self.txt_arren_rut, arrendador.get("rut"))
                set_text(self.txt_arren_direccion, arrendador.get("direccion"))
                set_text(self.txt_arren_telefono, arrendador.get("telefono"))
                set_text(self.txt_arren_email, arrendador.get("correo_electronico"))

            # ===============================
            # ARRENDATARIO
            # ===============================
            arrendatario = self.detalle_arriendo.get('arrendatario', {})
            if isinstance(arrendatario, dict):
                set_text(self.txt_arrendatario_rut, arrendatario.get("rut"))
                set_text(self.txt_arrendatario_nombre, arrendatario.get("nombre"))
                set_text(self.txt_arrendatario_telefono, arrendatario.get("telefono"))
                set_text(self.txt_arrendatario_email, arrendatario.get("email"))
                set_text(self.txt_arrendatario_direccion, arrendatario.get("direccion"))

                self.cmb_tipo_trabajador.setCurrentText(
                    arrendatario.get("tipo_trabajador", "Dependiente")
                )

                set_text(self.txt_antiguedad_laboral, arrendatario.get("antiguedad_laboral"))
                self.cmb_dicom.setCurrentText(arrendatario.get("dicom", "Sí"))
                set_text(self.txt_comentarios, arrendatario.get("comentarios"))

                set_money(self.txt_renta, arrendatario.get("renta"))

            # ===============================
            # PROPIEDAD
            # ===============================
            propiedad = self.detalle_arriendo.get('propiedad', {})
            if isinstance(propiedad, dict):
                set_text(self.txt_prop_codigo, propiedad.get("codigo_interno"))
                set_text(self.txt_prop_direccion, propiedad.get("direccion"))
                set_text(self.txt_prop_rol, propiedad.get("rol"))
                set_text(self.txt_prop_comuna, propiedad.get("comuna"))
                self.cmb_dominio_vigente.setCurrentText(
                    propiedad.get("dominio_vigente", "No Posee Documento")
                )

            # ===============================
            # ARRIENDO
            # ===============================
            arriendo = self.detalle_arriendo.get('arriendo', {})
            if isinstance(arriendo, dict):

                set_date(self.fecha_inicio, arriendo.get("fecha_inicio"))
                set_date(self.fecha_termino, arriendo.get("fecha_termino"))

                set_money(self.txt_renta_mensual, arriendo.get("renta_mensual"))
                set_money(self.txt_garantia, arriendo.get("garantia"))

                self.chk_gastos_comunes.setChecked(
                    bool(arriendo.get("gastos_comunes_incluidos", False))
                )

                self.cmb_estado_arriendo.setCurrentText(
                    arriendo.get("estado", "Vigente")
                )

                self.cmb_tipo_contrato.setCurrentText(
                    arriendo.get("tipo_contrato", "Plazo Fijo")
                )

                set_text(self.txt_forma_pago, arriendo.get("forma_pago"))
                self.cmb_periodo_pago.setCurrentText(
                    arriendo.get("periodo_pago", "Mensual")
                )

                self.cmb_cuenta_fm.setCurrentText(
                    str(arriendo.get("cuenta_fm", "2"))
                )

                set_text(self.txt_nro_cuenta, arriendo.get("nro_cuenta"))
                set_text(self.txt_banco_destino, arriendo.get("banco_destino"))
                set_text(self.txt_tipo_cuenta, arriendo.get("tipo_cuenta"))
                set_text(self.txt_rut_para_deposito, arriendo.get("rut_para_deposito"))
                set_text(self.txt_titular_deposito, arriendo.get("titular_deposito"))
                set_text(self.txt_quien_deposita, arriendo.get("quien_deposita"))
                set_text(self.txt_correo_deposito, arriendo.get("correo_deposito"))

                set_text(self.txt_aseo_municipal, arriendo.get("aseo_municipal"))
                set_money(self.txt_reajuste, arriendo.get("reajuste"))
                set_text(self.txt_ggcc, arriendo.get("ggcc"))
                set_text(self.txt_cuenta_ggcc, arriendo.get("cuenta_ggcc"))

                set_text(
                    self.txt_honorarios_porcentaje,
                    self.percent(arriendo.get("honorarios_porcentaje"))
                )

                set_money(self.txt_honorarios_monto, arriendo.get("honorarios_monto"))

                set_text(self.txt_dia_pago, arriendo.get("dia_pago"))
                set_text(self.txt_ultimo_mes_pago, arriendo.get("ultimo_mes_pago"))
                set_text(self.txt_direccion_consulta, arriendo.get("direccion_consulta"))
                set_text(self.txt_periodo_anterior, arriendo.get("periodo_anterior"))

                set_money(self.txt_monto_anterior, arriendo.get("monto_anterior"))

                set_text(self.txt_naturaleza_bien_raiz, arriendo.get("naturaleza_bien_raiz"))

                # 🔴 CORRECCIÓN IMPORTANTE
                self.cmb_dfl12.setCurrentText(
                    arriendo.get("dfl2", "Si")
                )

                set_text(self.txt_destino, arriendo.get("destino"))
                self.cmb_amoblado.setCurrentText(
                    arriendo.get("amoblado", "Si")
                )

                set_text(self.txt_observaciones, arriendo.get("observaciones"))

        except Exception as e:
            print(f"Error al cargar el arriendo: {e}")
            QMessageBox.warning(
                self,
                "Error",
                f"No se pudieron cargar los datos del arriendo:\n{e}"
            )




    def cargar_resultado_final_evaluacion(self, eva):
        """
        evaluacion: dict que viene desde la BD
        Solo setea la columna Resultado (4).
        """

        mapping = {
            0: eva.get("sueldo_base"),
            1: eva.get("gratificacion"),
            2: eva.get("total_imponible"),
            3: eva.get("locomocion"),
            4: eva.get("total_no_imponible"),
            5: eva.get("total_haberes"),
            6: eva.get("descuentos_legales"),
            7: eva.get("desc_varios"),
            8: eva.get("anticipo"),
            9: eva.get("liquido_pago"),
        }

        for fila, valor in mapping.items():
            if valor is None:
                continue

            item = QTableWidgetItem(f"{valor:,.0f}")
            item.setTextAlignment(Qt.AlignCenter)
            item.setFlags(Qt.ItemIsEnabled)  # solo lectura
            self.tbl_evaluacion.setItem(fila, 4, item)

    def get_int(self, value):
        if value is None:
            return 0.0

        if isinstance(value, (int, float)):
            return float(value)

        texto = str(value).strip()
        if not texto:
            return 0.0

        # eliminar símbolos
        texto = texto.replace("$", "").replace("%", "")

        # normalizar separadores
        # miles: .  → eliminar
        # decimales: , → .
        texto = texto.replace(".", "").replace(",", ".")

        try:
            return float(texto)
        except ValueError:
            return 0.0

    def obtener_diccionario_evaluacion_independiente(self):
        meses_map = {
            0: "01", 1: "02", 2: "03", 3: "04",
            4: "05", 5: "06", 6: "07", 7: "08",
            8: "09", 9: "10", 10: "11", 11: "12"
        }

        detalle = []
        total_ventas = 0
        total_compras = 0

        hoy = date.today()
        year = hoy.year

        for row in range(12):
            try:
                iva_debito = float(
                    (self.tbl_independiente.item(row, 1).text() or "0")
                    .replace(".", "").replace(",", ".")
                )
                iva_credito = float(
                    (self.tbl_independiente.item(row, 2).text() or "0")
                    .replace(".", "").replace(",", ".")
                )
            except (ValueError, AttributeError):
                continue

            # ✅ CÁLCULO REAL
            ventas = iva_debito / IVA_FACTOR
            compras = iva_credito / IVA_FACTOR

            periodo = f"{year}-{meses_map[row]}-01"

            detalle.append({
                'periodo': periodo,
                'iva_debito': iva_debito,
                'iva_credito': iva_credito,
                'ventas_netas_estimadas': ventas,
                'compras_netas_estimadas': compras
            })

            total_ventas += ventas
            total_compras += compras

        excedente = total_ventas - total_compras
        renta_anual = excedente * CASTIGO_DEFAULT
        renta_mensual = renta_anual / 12

        return {
            'periodo_desde': hoy.isoformat(),
            'resumen': {
                'ventas_anuales': total_ventas,
                'compras_anuales': total_compras,
                'excedente_anual': excedente,
                'renta_anual_estimada': renta_anual,
                'renta_mensual_estimada': renta_mensual
            },
            'detalle': detalle
        }
    from PySide6.QtWidgets import QFileDialog

    def obtener_ruta_excel(self):

        rut_arrendatario = self.txt_arrendatario_rut.text()

        rut = rut_arrendatario.strip().replace(".", "").replace("-", "")
        nombre_archivo = f"evaluacion_independiente_{rut}.xlsx"

        ruta_archivo, _ = QFileDialog.getSaveFileName(
            self,
            "Guardar evaluación",
            nombre_archivo,
            "Excel (*.xlsx)"
        )

        if not ruta_archivo:
            return  # usuario canceló

        if not ruta_archivo.lower().endswith(".xlsx"):
            ruta_archivo += ".xlsx"

        # 👉 generar el excel
        self.generar_excel_evaluacion_independiente(ruta_archivo)


    def generar_excel_evaluacion_independiente(self, ruta_archivo):
        datos = self.obtener_diccionario_evaluacion_independiente()

        wb = Workbook()
        ws = wb.active
        ws.title = "Evaluación Independiente"

        # -----------------------------
        # ESTILOS
        # -----------------------------
        bold = Font(bold=True)

        thin = Side(style="thin")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        formato_contable = '"$"#,##0;("$"#,##0)'

        # -----------------------------
        # TITULO
        # -----------------------------
        ws["A1"] = "EVALUACIÓN TRABAJADOR INDEPENDIENTE"
        ws["A1"].font = bold
        ws.append([])

        # -----------------------------
        # TABLA MENSUAL
        # -----------------------------
        headers = [
            "Mes",
            "IVA Débito",
            "IVA Crédito",
            "Ventas Netas Estimadas",
            "Compras Netas Estimadas",
            "Excedente"
        ]

        ws.append(headers)
        header_row = ws.max_row

        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=header_row, column=col)
            cell.font = bold
            cell.border = border

        inicio_tabla = header_row + 1

        for detalle in datos["detalle"]:
            excedente = (
                detalle["ventas_netas_estimadas"]
                - detalle["compras_netas_estimadas"]
            )

            ws.append([
                detalle["periodo"][:7],
                detalle["iva_debito"],
                detalle["iva_credito"],
                detalle["ventas_netas_estimadas"],
                detalle["compras_netas_estimadas"],
                excedente
            ])

        fin_tabla = ws.max_row

        # Aplicar bordes y formato contable
        for row in range(inicio_tabla, fin_tabla + 1):
            for col in range(1, 7):
                cell = ws.cell(row=row, column=col)
                cell.border = border

                if col >= 2:  # columnas monetarias
                    cell.number_format = formato_contable

        # -----------------------------
        # RESUMEN ANUAL
        # -----------------------------
        ws.append([])
        ws.append(["RESUMEN ANUAL"])
        ws["A" + str(ws.max_row)].font = bold

        resumen = datos["resumen"]

        resumen_filas = [
            ("Ventas Anuales", resumen["ventas_anuales"]),
            ("Compras Anuales", resumen["compras_anuales"]),
            ("Excedente Anual", resumen["excedente_anual"]),
            ("Factor Castigo", f"{CASTIGO_DEFAULT * 100:.0f}%"),
            ("Renta Anual Estimada", resumen["renta_anual_estimada"]),
            ("Renta Mensual Estimada", resumen["renta_mensual_estimada"]),
        ]

        for etiqueta, valor in resumen_filas:
            ws.append([etiqueta, valor])
            fila = ws.max_row
            ws.cell(row=fila, column=1).font = bold

            if isinstance(valor, (int, float)):
                ws.cell(row=fila, column=2).number_format = formato_contable

        # -----------------------------
        # METODOLOGÍA
        # -----------------------------
        ws.append([])
        ws.append(["METODOLOGÍA DE CÁLCULO"])
        ws["A" + str(ws.max_row)].font = bold

        metodologia = [
            "IVA Débito: IVA declarado por ventas del período.",
            "IVA Crédito: IVA declarado por compras o gastos del período.",
            "Ventas Netas Estimadas = IVA Débito / 0,19.",
            "Compras Netas Estimadas = IVA Crédito / 0,19.",
            "Excedente Mensual = Ventas Netas - Compras Netas.",
            f"Renta Anual Estimada = Excedente Anual x Factor Castigo ({CASTIGO_DEFAULT}).",
            "Renta Mensual Estimada = Renta Anual / 12.",
            f"Fecha de evaluación: {date.today().isoformat()}",
        ]

        for linea in metodologia:
            ws.append([linea])

        # -----------------------------
        # AJUSTE DE ANCHO
        # -----------------------------
        ws.column_dimensions["A"].width = 22
        ws.column_dimensions["B"].width = 18
        ws.column_dimensions["C"].width = 18
        ws.column_dimensions["D"].width = 24
        ws.column_dimensions["E"].width = 26
        ws.column_dimensions["F"].width = 18

        wb.save(ruta_archivo)

    def get_porcentaje(self, value):
        if not value or value.strip() == "":
            return 0.0
        try:
            limpio = value.replace('%', '').replace(',', '.').strip()
            return float(limpio) / 100
        except:
            return 0.0


    def guardar_arriendo(self):
        try:
            valores_tabla = {}

            if self.cmb_tipo_trabajador.currentText() == "Dependiente":
                valores_tabla = self.obtener_diccionario_evaluacion()

            data_arriendo = {

                'propiedad': {
                    'codigo': self.txt_prop_codigo.text(),
                    'direccion': self.txt_prop_direccion.text(),
                    'rol': self.txt_prop_rol.text(),
                    'comuna': self.txt_prop_comuna.text(),
                    'dominio_vigente': self.cmb_dominio_vigente.currentText()
                },

                'arrendador': {
                    'rut': self.txt_arren_rut.text(),
                    'nombre': self.txt_arren_nombre.text(),
                    'direccion': self.txt_arren_direccion.text(),
                    'telefono': self.txt_arren_telefono.text(),
                    'correo_electronico': self.txt_arren_email.text()
                },

                'arrendatario': {
                    'rut': self.txt_arrendatario_rut.text(),
                    'nombre': self.txt_arrendatario_nombre.text(),
                    'telefono': self.txt_arrendatario_telefono.text(),
                    'email': self.txt_arrendatario_email.text(),
                    'direccion': self.txt_arrendatario_direccion.text(),
                    'tipo_trabajador': self.cmb_tipo_trabajador.currentText(),
                    'dicom': self.cmb_dicom.currentText(),
                    'comentarios': self.txt_comentarios.toPlainText(),
                    'evaluacion_estado': self.cmb_evaluacion_estado.currentText(),
                    'antiguedad_laboral': self.txt_antiguedad_laboral.text(),
                    'fecha_evaluacion': self.fecha_evaluacion.date().toString("yyyy-MM-dd"),
                    'renta': self.get_int(self.txt_renta.text())
                },

                'arriendo': {
                    'fecha_inicio': self.fecha_inicio.date().toString("yyyy-MM-dd"),
                    'fecha_termino': (
                            self.fecha_termino.date().toString("yyyy-MM-dd")
                            if self.cmb_estado_arriendo.currentText() != "Vigente"
                            else None
                        ),
                    'tipo_documento':self.cmb_tipo_documento.currentText(),
                    'renta_mensual': self.get_int(self.txt_renta_mensual.text()),
                    'garantia': self.get_int(self.txt_garantia.text()),
                    'gastos_comunes_incluidos': self.chk_gastos_comunes.isChecked(),
                    'estado': self.cmb_estado_arriendo.currentText(),
                    'tipo_contrato': self.cmb_tipo_contrato.currentText(),
                    'forma_pago': self.txt_forma_pago.text(),
                    'periodo_pago': self.cmb_periodo_pago.currentText(),
                    'cuenta_fm': self.cmb_cuenta_fm.currentText(),

                    'nro_cuenta': self.txt_nro_cuenta.text(),
                    'banco_destino': self.txt_banco_destino.text(),
                    'aseo_municipal': self.txt_aseo_municipal.text(),
                    'reajuste': self.txt_reajuste.text(),
                    'ggcc': self.txt_ggcc.text(),

                    'honorarios_porcentaje': self.get_porcentaje(self.txt_honorarios_porcentaje.text()),
                    'honorarios_monto': self.get_int(self.txt_honorarios_monto.text()),
                    'titular_deposito': self.txt_titular_deposito.text(),
                    'quien_deposita': self.txt_quien_deposita.text(),
                    'correo_deposito': self.txt_correo_deposito.text(),
                    'dia_pago': self.txt_dia_pago.text(),
                    'tipo_cuenta': self.txt_tipo_cuenta.text(),
                    'rut_para_deposito': self.txt_rut_para_deposito.text(),
                    'cuenta_ggcc': self.txt_cuenta_ggcc.text(),
                    'ultimo_mes_pago': self.txt_ultimo_mes_pago.text(),
                    'direccion_consulta': self.txt_direccion_consulta.text(),
                    'periodo_anterior': self.txt_periodo_anterior.text(),
                    'monto_anterior': self.get_int(self.txt_monto_anterior.text()),
                    'naturaleza_bien_raiz': self.txt_naturaleza_bien_raiz.text(),
                    'tipo_documento': self.cmb_tipo_documento.currentText(),
                    'dfl2': self.cmb_dfl12.currentText(),
                    'destino': self.txt_destino.text(),
                    'amoblado': self.cmb_amoblado.currentText(),
                    'observaciones': self.txt_observaciones.toPlainText()
                }
            }

            if self.cmb_tipo_trabajador.currentText() == "Dependiente":
                data_arriendo['evaluacion_arrendatario'] = {
                    'fecha_evaluacion': self.fecha_evaluacion_arrendatario.date().toString("yyyy-MM-dd"),
                    'sueldo_base': self.get_int(valores_tabla.get('sueldo_base', 0)),
                    'gratificacion': self.get_int(valores_tabla.get('gratificacion', 0)),
                    'total_imponible': self.get_int(valores_tabla.get('total_imponible', 0)),
                    'total_no_imponible': self.get_int(valores_tabla.get('total_no_imponible', 0)),
                    'descuentos_legales': self.get_int(valores_tabla.get('descuentos_legales', 0)),
                    'liquido_pago': self.get_int(valores_tabla.get('liquido_pago', 0)),
                    'anticipo': self.get_int(valores_tabla.get('anticipo', 0)),
                    'desc_varios': self.get_int(valores_tabla.get('desc_varios', 0)),
                    'locomocion': self.get_int(valores_tabla.get('locomocion', 0)),
                    'im_renta': self.get_int(self.txt_impuesto_renta.text())
                }

            if self.cmb_tipo_trabajador.currentText() == "Independiente":
                evaluacion = self.obtener_diccionario_evaluacion_independiente()

                data_arriendo['evaluacion_independiente'] = {
                    'rut_arrendatario': self.txt_arrendatario_rut.text(),
                    'periodo_desde': date.today().isoformat(),
                    'factor_castigo': CASTIGO_DEFAULT,
                    

                    # ⬇️ SOLO RESUMEN
                    **evaluacion['resumen']
                }

                # ⬇️ DETALLE VA APARTE
                data_arriendo['evaluacion_independiente_detalle'] = evaluacion['detalle']


            arriendo_id = guardar_arriendo(data_arriendo)

            if arriendo_id:
                QMessageBox.information(
                    self,
                    "Éxito",
                    f"El arriendo se ha guardado correctamente.\n\nID: {arriendo_id}"
                )
                self.arriendo_guardado.emit()
                self.close()

        except Exception as e:
            QMessageBox.critical(self, "Error", f"No se pudo guardar el arriendo: {str(e)}")


            

            


    def crear_label(self, texto, obligatorio=False):
        label = QLabel(texto)
        if obligatorio:
            label.setProperty("obligatorio", "true")
        return label
    

    def validar_rut(self, rut):
        # Implementación básica de validación de RUT chileno
        rut = rut.replace(".", "").replace("-", "").upper()
        if not rut[:-1].isdigit():
            return False
        
        cuerpo = rut[:-1]
        dv = rut[-1]
        
        suma = 0
        multiplicador = 2
        
        for c in reversed(cuerpo):
            suma += int(c) * multiplicador
            multiplicador += 1
            if multiplicador > 7:
                multiplicador = 2
        
        resto = suma % 11

        dv_calculado = 11-resto

        if dv_calculado == 11:
            dv_calculado = "0"
        elif dv_calculado == 10:
            dv_calculado = "K"
        else:
            dv_calculado = str(dv_calculado)
        
        return dv == dv_calculado
    

    def validar_campos_obligatorios(self):
        # Validar campos obligatorios básicos
        campos_obligatorios = [
            (self.txt_prop_codigo.text(), "Código de propiedad"),
            (self.txt_prop_direccion.text(), "Dirección de propiedad"),
            (self.txt_prop_rol.text(), "ROL de propiedad"),
            (self.txt_prop_comuna.text(), "Comuna de propiedad")
        ]
        
        

        campos_obligatorios += [
                (self.txt_arren_nombre.text(), "Nombre del arrendador"),
                (self.txt_arren_rut.text(), "RUT del arrendador"),
                (self.txt_arrendatario_nombre.text(), "Nombre del arrendatario"),
                (self.txt_arrendatario_rut.text(), "RUT del arrendatario"),
        ]


        for valor, nombre in campos_obligatorios:
            if not valor.strip():
                QMessageBox.warning(self, "Campo obligatorio", f"El campo {nombre} es obligatorio")
                return False
        
        # Validar RUTs
        
        if not self.validar_rut(self.txt_arren_rut.text()):
            QMessageBox.warning(self, "RUT inválido", "El RUT del arrendador no es válido")
            return False
            
        if not self.validar_rut(self.txt_arrendatario_rut.text()):
            QMessageBox.warning(self, "RUT inválido", "El RUT del arrendatario no es válido")
            return False
    

class FormularioMes(QDialog):
    def __init__(self, numero_mes, parent=None, rut_arrendatario=None):
        super().__init__(parent)
        self.setWindowTitle(f"Evaluación - Mes {numero_mes}")
        self.resize(650, 600)

        # --- LAYOUT PRINCIPAL ---
        layout_principal = QVBoxLayout(self)

        # --- SCROLL AREA ---
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        layout_principal.addWidget(scroll)

        self.rut_arrendatario = rut_arrendatario

        # --- CONTENEDOR INTERNO (donde va todo tu contenido) ---
        contenedor = QWidget()
        scroll.setWidget(contenedor)

        layout = QVBoxLayout(contenedor)

        form = QFormLayout()

        # ---- CAMPOS ----
        self.sueldo_base = QLineEdit()
        self.horas_extras = QLineEdit()
        self.afp = QLineEdit()
        self.salud = QLineEdit()
        self.cesantia = QLineEdit()
        self.bono1 = QLineEdit()
        self.bono2 = QLineEdit()
        self.bono3 = QLineEdit()
        self.colacion = QLineEdit()
        self.locomocion = QLineEdit()
        self.cargas_familiares = QLineEdit()
        self.viaticos = QLineEdit()
        self.herramientas = QLineEdit()
        self.anticipo = QLineEdit()
        self.otros = QLineEdit()
        self.sueldo_minimo = QLineEdit()

        form.addRow("Sueldo Base:", self.sueldo_base)
        form.addRow("Horas Extras:", self.horas_extras)
        form.addRow("AFP:", self.afp)
        form.addRow("Salud:", self.salud)
        form.addRow("Cesantía:", self.cesantia)
        form.addRow("Bono 1:", self.bono1)
        form.addRow("Bono 2:", self.bono2)
        form.addRow("Bono 3:", self.bono3)
        form.addRow("Colación:", self.colacion)
        form.addRow("Locomoción:", self.locomocion)
        form.addRow("Cargas Familiares:", self.cargas_familiares)
        form.addRow("Anticipo:", self.anticipo)
        form.addRow("Viáticos:", self.viaticos)
        form.addRow("Herramientas:", self.herramientas)
        form.addRow("Sueldo Mínimo:", self.sueldo_minimo)

        layout.addLayout(form)

        # ---- TABLA DE OTROS DESCUENTOS ----
        layout.addWidget(QLabel("Otros Descuentos:"))

        self.tbl_descuentos = QTableWidget()
        self.tbl_descuentos.setColumnCount(2)
        self.tbl_descuentos.setHorizontalHeaderLabels(["Concepto", "Monto"])
        self.tbl_descuentos.horizontalHeader().setStretchLastSection(True)

        self.tbl_descuentos.setMinimumHeight(120)

        layout.addWidget(self.tbl_descuentos)

        # Botones descuentos
        btns = QHBoxLayout()
        btn_add = DoubleClickButton("Agregar descuento")
        btn_del = DoubleClickButton("Eliminar seleccionado")

        btn_add.doubleClicked.connect(self.agregar_descuento)
        btn_del.doubleClicked.connect(self.eliminar_descuento)

        btns.addWidget(btn_add)
        btns.addWidget(btn_del)
        layout.addLayout(btns)

        # Botones OK/Cancel
        botones = QHBoxLayout()
        btn_ok = DoubleClickButton("Aceptar")
        btn_cancel = DoubleClickButton("Cancelar")
        btn_excel = DoubleClickButton("Exportar a Excel")

        btn_ok.doubleClicked.connect(self.aceptar)
        btn_cancel.doubleClicked.connect(self.reject)
        btn_excel.doubleClicked.connect(self.exportar_evaluacion_excel)

        botones.addWidget(btn_ok)
        botones.addWidget(btn_cancel)
        botones.addWidget(btn_excel)
        layout.addLayout(botones)



    # ---------- Obtener datos para procesarlos ----------
    def obtener_datos(self):
        descuentos = []
        for i in range(self.tbl_descuentos.rowCount()):
            item_concepto = self.tbl_descuentos.item(i, 0)
            item_monto = self.tbl_descuentos.item(i, 1)

            if item_concepto is None or item_monto is None:
                continue

            concepto = item_concepto.text().strip()
            if not concepto:
                continue

            try:
                monto = float(item_monto.text() or 0)
            except:
                monto = 0

            descuentos.append({"nombre": concepto, "monto": monto})

        return {
            "sueldo_base": float(self.sueldo_base.text() or 0),
            "horas_extras": float(self.horas_extras.text() or 0),
            "afp": float(self.afp.text() or 0),
            "salud": float(self.salud.text() or 0),
            "cesantia": float(self.cesantia.text() or 0),
            "bono1": float(self.bono1.text() or 0),
            "bono2": float(self.bono2.text() or 0),
            "bono3": float(self.bono3.text() or 0),
            "colacion": float(self.colacion.text() or 0),
            "locomocion": float(self.locomocion.text() or 0),
            "viaticos": float(self.viaticos.text() or 0),
            "herramientas": float(self.herramientas.text() or 0),
            "otros": float(self.otros.text() or 0) if hasattr(self, "otros") else 0,
            "anticipo": float(self.anticipo.text() or 0) if hasattr(self, "anticipo") else 0,
            "cargas_familiares": float(self.cargas_familiares.text() or 0),

            # ⚠ IMPORTANTE: debes tener un input para sueldo mínimo
            "sueldo_minimo": float(self.sueldo_minimo.text() or 0),

            # Descuentos varios
            "descuentos_varios": descuentos
        }

    
    def agregar_descuento(self):
        row = self.tbl_descuentos.rowCount()
        self.tbl_descuentos.insertRow(row)
        self.tbl_descuentos.setItem(row, 0, QTableWidgetItem(""))
        self.tbl_descuentos.setItem(row, 1, QTableWidgetItem("0"))

    # ---------- Función para eliminar un descuento ----------
    def eliminar_descuento(self):
        row = self.tbl_descuentos.currentRow()
        if row >= 0:
            self.tbl_descuentos.removeRow(row)

    # ---------- Obtener datos para procesarlos ----------
    def calcular_totales_mes(self, datos):

        # 1) Tope legal gratificación
        tope_grat = datos["sueldo_minimo"] * 4.75 / 12

        # 2) Gratificación (25% SB topado)
        gratificacion = min(datos["sueldo_base"] * 0.25, tope_grat)

        # 3) Total imponible
        th_imponibles = (
            datos["sueldo_base"]
            + gratificacion
            + datos["bono1"]
            + datos["bono2"]
            + datos["bono3"]
            + datos["horas_extras"]
        )

        # 4) Total no imponible
        th_no_imponibles = (
            datos["colacion"]
            + datos["locomocion"]
            + datos["viaticos"]
            + datos["herramientas"]
            + datos['cargas_familiares']
            + datos["otros"]
        )

        # 5) Total haberes
        total_haberes = th_imponibles + th_no_imponibles

        # 6) Descuentos legales
        descuentos_legales = datos["afp"] + datos["salud"] + datos["cesantia"]

        # 7) Descuentos varios
        descuentos_varios_total = sum(d["monto"] for d in datos.get("descuentos_varios", []))

        # 8) Anticipo
        anticipo = datos.get("anticipo", 0)

        # 9) Líquido a pago
        liquido = total_haberes - descuentos_legales - descuentos_varios_total - anticipo


        return {
            "tope_gratificacion": tope_grat,
            "gratificacion": gratificacion,
            "th_imponibles": th_imponibles,
            "th_no_imponibles": th_no_imponibles,
            "total_haberes": total_haberes,
            "descuentos_legales": descuentos_legales,
            "descuentos_varios": descuentos_varios_total,
            "anticipo": anticipo,
            "liquido": liquido
        }
    

    def aceptar(self):
        datos = self.obtener_datos()
        resultados = self.calcular_totales_mes(datos)

        # Guardamos ambos: datos originales + cálculos
        self.resultado_final = {
            "datos": datos,
            "calculos": resultados
        }

        self.accept()

    def exportar_evaluacion_excel(self):
        datos = self.obtener_datos()
        calculos = self.calcular_totales_mes(datos)

        self.resultado_final = {
            "datos": datos,
            "calculos": calculos
        }

        rut = self.rut_arrendatario.strip().replace(".", "").replace("-", "")
        nombre_archivo = f"evaluacion_arrendatario_{rut}.xlsx"

        ruta, _ = QFileDialog.getSaveFileName(
            self,
            "Guardar evaluación",
            nombre_archivo,
            "Excel (*.xlsx)"
        )

        if not ruta:
            return

        wb = Workbook()
        ws = wb.active
        ws.title = "Evaluación"

        bold = Font(bold=True)

        fila = 1

        # =========================
        # DATOS DE ENTRADA
        # =========================
        ws[f"A{fila}"] = "DATOS DE ENTRADA"
        ws[f"A{fila}"].font = bold
        fila += 2

        for clave, valor in datos.items():
            if clave == "descuentos_varios":
                continue
            ws[f"A{fila}"] = clave.replace("_", " ").title()
            ws[f"B{fila}"] = valor
            fila += 1

        # Descuentos varios
        fila += 1
        ws[f"A{fila}"] = "Descuentos Varios"
        ws[f"A{fila}"].font = bold
        fila += 1

        for d in datos.get("descuentos_varios", []):
            ws[f"A{fila}"] = d["nombre"]
            ws[f"B{fila}"] = d["monto"]
            fila += 1

        # Margen grande
        fila += 3

        # =========================
        # CÁLCULOS
        # =========================
        ws[f"A{fila}"] = "CÁLCULOS"
        ws[f"A{fila}"].font = bold
        fila += 2

        for clave, valor in calculos.items():
            ws[f"A{fila}"] = clave.replace("_", " ").title()
            ws[f"B{fila}"] = valor
            fila += 1

        fila += 3

        # =========================
        # EXPLICACIÓN DE CÁLCULOS
        # =========================
        ws[f"A{fila}"] = "EXPLICACIÓN DE LOS CÁLCULOS"
        ws[f"A{fila}"].font = bold
        fila += 2

        explicaciones = [
            "1) Tope de gratificación: sueldo mínimo × 4.75 / 12.",
            "2) Gratificación: 25% del sueldo base, con tope legal.",
            "3) Total imponible: sueldo base + gratificación + bonos + horas extras.",
            "4) Total no imponible: colación, locomoción, viáticos, herramientas, cargas familiares y otros.",
            "5) Total haberes: imponibles + no imponibles.",
            "6) Descuentos legales: AFP + salud + cesantía.",
            "7) Descuentos varios: suma de descuentos adicionales ingresados.",
            "8) Anticipo: monto adelantado al trabajador.",
            "9) Líquido a pago: total haberes − descuentos − anticipo."
        ]

        for texto in explicaciones:
            ws[f"A{fila}"] = texto
            fila += 1

        # =========================
        # AJUSTES VISUALES
        # =========================
        ws.column_dimensions["A"].width = 45
        ws.column_dimensions["B"].width = 20

        try:
            wb.save(ruta)
            QMessageBox.information(self, "Éxito", f"Archivo exportado correctamente:\n{ruta}")
        except Exception as e:
            QMessageBox.critical(self, "Error", f"No se pudo guardar el archivo:\n{e}")

    
class ImpuestoRentaDialog(QDialog):
    def __init__(self, renta_tributable, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Calcular Impuesto a la Renta")

        self.renta_tributable = renta_tributable
        self.resultado = None  # ← aquí guardaremos el cálculo final

        layout = QFormLayout()

        # Campo RT (solo lectura)
        self.txt_rt = QLineEdit(str(renta_tributable))
        self.txt_rt.setReadOnly(True)
        layout.addRow("Renta Tributable:", self.txt_rt)

        # Campos editables
        self.txt_multiplicador = QLineEdit()
        layout.addRow("Tasa:", self.txt_multiplicador)

        self.txt_resta = QLineEdit()
        layout.addRow("Rebaja:", self.txt_resta)

        # Botones
        btn_calcular = DoubleClickButton("Calcular")
        btn_cancelar = DoubleClickButton("Cancelar")

        btn_calcular.doubleClicked.connect(self.calcular)
        btn_cancelar.doubleClicked.connect(self.reject)

        hbox = QHBoxLayout()
        hbox.addWidget(btn_calcular)
        hbox.addWidget(btn_cancelar)

        layout.addRow(hbox)
        self.setLayout(layout)

    def calcular(self):
        try:
            rt = float(self.renta_tributable)
            multiplicador = float(self.txt_multiplicador.text())
            resta = float(self.txt_resta.text())

            self.resultado = (rt * multiplicador) - resta
            self.accept()

        except ValueError:
            pass  # puedes mostrar un QMessageBox si quieres


class DialogFactorActualizacion(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Factor de actualización de arriendos")
        self.resize(500, 400)

        self.resultado = []

        layout = QVBoxLayout(self)

        self.tabla = QTableWidget(12, 2)
        self.tabla.setHorizontalHeaderLabels([
            "Mes",
            "Factor de actualización"
        ])

        meses = [
            "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
            "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"
        ]

        for row, mes in enumerate(meses):
            item_mes = QTableWidgetItem(mes)
            item_mes.setFlags(Qt.ItemIsEnabled)
            self.tabla.setItem(row, 0, item_mes)

            factor = QDoubleSpinBox()
            factor.setDecimals(3)
            factor.setRange(0.5, 5)
            factor.setSingleStep(0.001)
            factor.setValue(1.0)
            self.tabla.setCellWidget(row, 1, factor)

        self.tabla.resizeColumnsToContents()
        layout.addWidget(self.tabla)

        botones = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        botones.accepted.connect(self.guardar)
        botones.rejected.connect(self.reject)
        layout.addWidget(botones)

    def guardar(self):
        self.resultado.clear()

        for row in range(12):
            mes = self.tabla.item(row, 0).text()
            factor = self.tabla.cellWidget(row, 1).value()

            self.resultado.append({
                "mes": mes,
                "factor": factor
            })

        self.accept()



class DialogAbonosArriendo(QDialog):
    def __init__(self, arriendo_id, parent=None):
        super().__init__(parent)
        self.arriendo_id = arriendo_id
        self.setWindowTitle("Abonos del arriendo")

        layout = QVBoxLayout(self)

        self.tabla = QTableWidget(0, 3)
        self.tabla.setHorizontalHeaderLabels([
            "Fecha",
            "Monto",
            "Descripción"
        ])
        self.tabla.horizontalHeader().setStretchLastSection(True)

        layout.addWidget(self.tabla)

        btn_agregar = QPushButton("Agregar abono")
        btn_eliminar = QPushButton("Eliminar abono")
        btn_guardar = QPushButton("Guardar abonos")

        btn_agregar.clicked.connect(self.agregar_fila)
        btn_eliminar.clicked.connect(self.eliminar_fila)
        btn_guardar.clicked.connect(self.guardar_abonos)

        botones = QHBoxLayout()
        botones.addWidget(btn_agregar)
        botones.addWidget(btn_eliminar)
        botones.addStretch()
        botones.addWidget(btn_guardar)

        layout.addLayout(botones)

    def agregar_fila(self):
        fila = self.tabla.rowCount()
        self.tabla.insertRow(fila)

        self.tabla.setItem(fila, 0, QTableWidgetItem(QDate.currentDate().toString("yyyy-MM-dd")))
        self.tabla.setItem(fila, 1, QTableWidgetItem(""))
        self.tabla.setItem(fila, 2, QTableWidgetItem(""))


    def eliminar_fila(self):
        fila = self.tabla.currentRow()
        if fila >= 0:
            self.tabla.removeRow(fila)


    def guardar_abonos(self):
        abonos = []

        for row in range(self.tabla.rowCount()):
            fecha = self.tabla.item(row, 0).text()
            monto = self.tabla.item(row, 1).text()
            descripcion = self.tabla.item(row, 2).text()

            if not monto:
                continue

            abonos.append({
                "fecha": fecha,
                "monto": float(monto),
                "descripcion": descripcion
            })

        try:
            guardar_abonos_arriendo(self.arriendo_id, abonos)
            QMessageBox.information(self, "OK", "Abonos guardados correctamente")
            self.accept()
        except Exception as e:
            QMessageBox.critical(self, "Error", str(e))



class DialogGastosArriendo(QDialog):
    def __init__(self, arriendo_id, parent=None):
        super().__init__(parent)
        self.arriendo_id = arriendo_id
        self.setWindowTitle("Gastos del arriendo")

        layout = QVBoxLayout(self)

        self.tabla = QTableWidget(0, 3)
        self.tabla.setHorizontalHeaderLabels([
            "Fecha",
            "Monto",
            "Descripción"
        ])
        self.tabla.horizontalHeader().setStretchLastSection(True)

        layout.addWidget(self.tabla)

        btn_agregar = DoubleClickButton("Agregar gasto")
        btn_eliminar = DoubleClickButton("Eliminar gasto")
        btn_guardar = DoubleClickButton("Guardar gastos")

        btn_agregar.doubleClicked.connect(self.agregar_fila)
        btn_eliminar.doubleClicked.connect(self.eliminar_fila)
        btn_guardar.doubleClicked.connect(self.guardar_gastos)

        botones = QHBoxLayout()
        botones.addWidget(btn_agregar)
        botones.addWidget(btn_eliminar)
        botones.addStretch()
        botones.addWidget(btn_guardar)

        layout.addLayout(botones)

    def agregar_fila(self):
        fila = self.tabla.rowCount()
        self.tabla.insertRow(fila)

        self.tabla.setItem(fila, 0, QTableWidgetItem(QDate.currentDate().toString("yyyy-MM-dd")))
        self.tabla.setItem(fila, 1, QTableWidgetItem(""))
        self.tabla.setItem(fila, 2, QTableWidgetItem(""))


    def eliminar_fila(self):
        fila = self.tabla.currentRow()
        if fila >= 0:
            self.tabla.removeRow(fila)


    def guardar_gastos(self):
        gastos = []

        for row in range(self.tabla.rowCount()):
            fecha = self.tabla.item(row, 0).text()
            monto = self.tabla.item(row, 1).text()
            descripcion = self.tabla.item(row, 2).text()

            if not monto:
                continue

            gastos.append({
                "fecha": fecha,
                "monto": float(monto),
                "descripcion": descripcion
            })

        try:
            guardar_gastos_arriendo(self.arriendo_id, gastos)
            QMessageBox.information(self, "OK", "Gastos guardados correctamente")
            self.accept()
        except Exception as e:
            QMessageBox.critical(self, "Error", str(e))
